"""Seed stub templates from the GotIt Legal Sales catalog xlsx.

The catalog is the canonical list of 17 approved B2B Sales templates × up to
2 language variants. Each row in the xlsx becomes 1 stub Document template
(or 2 — one VIE, one BIL) with a small DOCX skeleton + placeholder fields.

Use case: demo / smoke testing. The DOCX content is a placeholder shell — the
real Got It DOCX files (referenced by file_code in the catalog) should be
uploaded later via the /templates UI to replace the stubs.

Run:
    python scripts/seed_legal_catalog.py <path/to/catalog.xlsx>
    python scripts/seed_legal_catalog.py <path/to/catalog.xlsx> --force
"""
from __future__ import annotations

import argparse
import asyncio
import io
import sys
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).parent.parent))

from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.orm import sessionmaker

from src.core.config import get_settings
from src.models.document import Document
from src.models.storage import StorageConfig
from src.models.user import User
from src.services.legal_catalog_parser import (
    CatalogEntry,
    categories_in_order,
    parse_legal_catalog,
)
from src.services.storage import get_storage_backend

ADMIN_USER_ID = uuid.UUID("00000000-0000-0000-0000-000000000002")
SEED_MARKER_KEY = "legal_catalog_seed"


def _force_utf8_stdout() -> None:
    if hasattr(sys.stdout, "buffer"):
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")


# Common placeholder fields every Sales B2B template needs. Granular sections
# come later when the real DOCX is uploaded and re-extracted.
COMMON_FIELDS: list[dict[str, Any]] = [
    {
        "placeholder": "ten_khach_hang",
        "label": "Tên khách hàng",
        "type": "blank",
        "required": True,
        "section_key": "customer",
    },
    {
        "placeholder": "ma_so_thue",
        "label": "Mã số thuế",
        "type": "blank",
        "required": True,
        "section_key": "customer",
    },
    {
        "placeholder": "dia_chi",
        "label": "Địa chỉ ĐKKD",
        "type": "blank",
        "required": True,
        "section_key": "customer",
    },
    {
        "placeholder": "nguoi_dai_dien",
        "label": "Người đại diện",
        "type": "blank",
        "required": True,
        "section_key": "customer",
    },
    {
        "placeholder": "chuc_vu",
        "label": "Chức vụ",
        "type": "blank",
        "required": True,
        "section_key": "customer",
    },
    {
        "placeholder": "so_hop_dong",
        "label": "Số hợp đồng",
        "type": "blank",
        "required": False,
        "section_key": "contract",
    },
    {
        "placeholder": "ngay_ky",
        "label": "Ngày ký",
        "type": "date",
        "required": True,
        "section_key": "contract",
    },
    {
        "placeholder": "ngay_hieu_luc",
        "label": "Ngày hiệu lực",
        "type": "date",
        "required": False,
        "section_key": "contract",
    },
    {
        "placeholder": "ngay_het_han",
        "label": "Ngày hết hạn",
        "type": "date",
        "required": False,
        "section_key": "contract",
    },
]

SECTIONS = [
    {"key": "customer", "label": "Thông tin Khách hàng", "order": 1},
    {"key": "contract", "label": "Thông tin Hợp đồng", "order": 2},
]


def _stub_docx(entry: CatalogEntry) -> bytes:
    """Skeleton DOCX with the entry's metadata + standard placeholders.

    The point is to give the app *something* to render and fill — the real
    legal-approved DOCX content is uploaded later via /templates.
    """
    doc = DocxDocument()
    doc.add_heading(entry.name, level=0)
    if entry.language == "bilingual":
        doc.add_paragraph("(Bản song ngữ Tiếng Việt — English / Bilingual VN-EN)")

    if entry.guidance:
        doc.add_paragraph(f"Hướng dẫn áp dụng: {entry.guidance}")
    doc.add_paragraph(
        f"Mã văn bản (Legal): {entry.file_code}. "
        "Đây là bản stub seed từ catalog — upload file DOCX gốc qua "
        "/templates để thay thế nội dung."
    )

    doc.add_heading("BÊN A — Khách hàng", level=2)
    doc.add_paragraph("Tên: {ten_khach_hang}")
    doc.add_paragraph("Mã số thuế: {ma_so_thue}")
    doc.add_paragraph("Địa chỉ: {dia_chi}")
    doc.add_paragraph("Đại diện: {nguoi_dai_dien} — Chức vụ: {chuc_vu}")

    doc.add_heading("Thông tin văn bản", level=2)
    doc.add_paragraph("Số hợp đồng: {so_hop_dong}")
    doc.add_paragraph("Ngày ký: {ngay_ky}")
    doc.add_paragraph("Hiệu lực từ: {ngay_hieu_luc} đến {ngay_het_han}")

    if entry.language == "bilingual":
        doc.add_heading("PARTY A — Customer (English)", level=2)
        doc.add_paragraph(
            "Name: {ten_khach_hang}. Tax code: {ma_so_thue}. "
            "Address: {dia_chi}. Representative: {nguoi_dai_dien} — "
            "Title: {chuc_vu}."
        )

    doc.add_heading("ĐẠI DIỆN CÁC BÊN", level=2)
    doc.add_paragraph("BÊN A                                   BÊN B (Got It)")
    doc.add_paragraph("(Ký, đóng dấu, ghi rõ họ tên)           (Ký, đóng dấu, ghi rõ họ tên)")

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _build_template_fields() -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for idx, f in enumerate(COMMON_FIELDS):
        out.append(
            {
                "id": f"f{idx}",
                "placeholder": f["placeholder"],
                "label": f["label"],
                "description": "",
                "location": f"para_{idx}",
                "type": f["type"],
                "options": None,
                "section_key": f["section_key"],
                "required": f["required"],
            }
        )
    return out


def _title(entry: CatalogEntry) -> str:
    suffix = " (VIE)" if entry.language == "single" else " (Song ngữ)"
    return f"[Legal] {entry.name}{suffix}"


async def _delete_seed_templates(db: AsyncSession, owner_id: uuid.UUID) -> int:
    """Hard-delete every Document previously seeded by this script."""
    rows = (
        await db.execute(
            select(Document).where(
                Document.owner_id == owner_id,
                Document.source_type == "template",
                Document.source_metadata[SEED_MARKER_KEY].astext == "true",
            )
        )
    ).scalars().all()

    if not rows:
        return 0

    for doc in rows:
        storage_cfg = await db.get(StorageConfig, doc.storage_config_id)
        if storage_cfg:
            backend = get_storage_backend(storage_cfg)
            try:
                await backend.delete(doc.file_path)
            except Exception:
                pass
        await db.delete(doc)

    await db.commit()
    return len(rows)


async def main(xlsx_path: Path, force: bool) -> None:
    settings = get_settings()
    engine = create_async_engine(settings.database_url, echo=False)
    Session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    xlsx_bytes = xlsx_path.read_bytes()
    entries = parse_legal_catalog(xlsx_bytes)
    print(f"Parsed {len(entries)} entries from catalog ({xlsx_path.name}).")
    print(f"Categories: {categories_in_order(entries)}")

    async with Session() as db:
        admin = await db.get(User, ADMIN_USER_ID)
        if not admin:
            raise SystemExit(
                "Admin user not found — run alembic migrations first."
            )

        existing_q = await db.execute(
            select(Document.id).where(
                Document.owner_id == ADMIN_USER_ID,
                Document.source_type == "template",
                Document.source_metadata[SEED_MARKER_KEY].astext == "true",
            )
        )
        existing = list(existing_q.scalars().all())

        if existing and not force:
            print(
                f"Đã có {len(existing)} legal-catalog stub đã seed trước đó. "
                "Thêm --force để xoá + reseed."
            )
            await engine.dispose()
            return

        if existing and force:
            removed = await _delete_seed_templates(db, ADMIN_USER_ID)
            print(f"Removed {removed} prior stub(s).")

        storage_cfg = (
            await db.execute(
                select(StorageConfig).where(StorageConfig.is_default.is_(True)).limit(1)
            )
        ).scalar_one_or_none()
        if not storage_cfg:
            raise SystemExit("No default storage config.")

        backend = get_storage_backend(storage_cfg)
        template_fields = _build_template_fields()

        seeded = 0
        for entry in entries:
            docx_bytes = _stub_docx(entry)
            filename = f"legal_{entry.file_code}.docx".replace(" ", "_")
            save_result = await backend.save(docx_bytes, filename)

            doc = Document(
                id=uuid.uuid4(),
                title=_title(entry),
                description=entry.guidance or None,
                file_name=save_result.file_name,
                original_filename=filename,
                file_path=save_result.file_path,
                file_size=save_result.file_size,
                mime_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
                extension=".docx",
                checksum=save_result.checksum,
                storage_config_id=storage_cfg.id,
                owner_id=ADMIN_USER_ID,
                source_type="template",
                source_metadata={
                    "template_fields": template_fields,
                    "extraction_status": "completed",
                    "language_mode": entry.language,
                    "sections": SECTIONS,
                    SEED_MARKER_KEY: True,
                    "legal_catalog": {
                        "stt": entry.stt,
                        "doc_category": entry.doc_category,
                        "file_code": entry.file_code,
                    },
                },
            )
            db.add(doc)
            seeded += 1

        await db.commit()
        print(f"Seeded {seeded} legal-catalog stubs into DB.")

    await engine.dispose()


if __name__ == "__main__":
    _force_utf8_stdout()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path, help="Path to the Legal catalog xlsx.")
    parser.add_argument(
        "--force", action="store_true",
        help="Delete prior seed and recreate.",
    )
    args = parser.parse_args()
    if not args.xlsx.is_file():
        raise SystemExit(f"File not found: {args.xlsx}")
    asyncio.run(main(args.xlsx, force=args.force))
