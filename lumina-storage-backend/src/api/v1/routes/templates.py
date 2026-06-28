import re
import uuid
from io import BytesIO
from pathlib import Path

from fastapi import APIRouter, Depends, Query, Request
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from src.api.deps import CurrentUser
from src.core.config import get_settings
from src.core.database import get_db
from src.models.document import Document
from src.schemas.template import (
    PaginatedTemplatesResponse,
    TemplateFieldResponse,
    TemplateFieldsUpdateRequest,
    TemplateResponse,
    TemplateSection,
    TemplateUpdateRequest,
)

router = APIRouter(prefix="/templates", tags=["templates"])


def _build_template_response(doc: Document) -> TemplateResponse:
    """Convert a Document with source_type='template' to TemplateResponse."""
    meta = doc.source_metadata or {}
    fields_raw = meta.get("template_fields", [])
    fields = [
        TemplateFieldResponse(
            id=f.get("id", ""),
            placeholder=f.get("placeholder", ""),
            label=f.get("label", ""),
            description=f.get("description", ""),
            location=f.get("location", ""),
            type=f.get("type", ""),
            options=f.get("options"),
            section_key=f.get("section_key"),
            required=f.get("required"),
        )
        for f in fields_raw
    ]
    sections_raw = meta.get("sections", []) or []
    sections = [
        TemplateSection(
            key=s.get("key", ""),
            label=s.get("label", ""),
            order=s.get("order", 0),
        )
        for s in sections_raw
        if s.get("key")
    ]
    language_mode = meta.get("language_mode") or "single"
    if language_mode not in ("single", "bilingual"):
        language_mode = "single"
    source_doc_id = meta.get("source_document_id")
    return TemplateResponse(
        id=doc.id,
        title=doc.title,
        description=doc.description,
        original_filename=doc.original_filename,
        source_document_id=uuid.UUID(source_doc_id) if source_doc_id else None,
        template_fields=fields,
        extraction_status=meta.get("extraction_status"),
        extraction_error=meta.get("extraction_error"),
        field_count=len(fields),
        language_mode=language_mode,
        sections=sections,
        created_at=doc.created_at,
        updated_at=doc.updated_at,
    )


@router.get("", response_model=PaginatedTemplatesResponse)
async def list_templates(
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
    q: str | None = Query(None, description="Search by title or description"),
    limit: int = Query(20, ge=1, le=100),
    offset: int = Query(0, ge=0),
):
    """List all templates owned by current user."""
    query = (
        select(Document)
        .where(
            Document.source_type == "template",
            Document.owner_id == current_user.id,
            Document.deleted_at.is_(None),
        )
    )
    if q:
        pattern = f"%{q}%"
        query = query.where(
            (Document.title.ilike(pattern)) | (Document.description.ilike(pattern))
        )
    query = query.order_by(Document.created_at.desc()).limit(limit + 1).offset(offset)

    result = await db.execute(query)
    rows = list(result.scalars().all())
    has_more = len(rows) > limit
    items = [_build_template_response(d) for d in rows[:limit]]
    return PaginatedTemplatesResponse(items=items, has_more=has_more)


@router.get("/{template_id}", response_model=TemplateResponse)
async def get_template(
    template_id: uuid.UUID,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Get template details including field list."""
    from fastapi import HTTPException

    doc = await db.get(Document, template_id)
    if not doc or doc.deleted_at is not None:
        raise HTTPException(404, "Template not found")

    # Normal case: already a template document
    if doc.source_type == "template":
        return _build_template_response(doc)

    meta = doc.source_metadata or {}
    extraction_status = meta.get("extraction_status")

    # Not a template and no extraction was triggered
    if not extraction_status:
        raise HTTPException(404, "Template not found")

    # Fast path: template_id was stored in meta by updated worker code
    if meta.get("template_id"):
        try:
            t = await db.get(Document, uuid.UUID(str(meta["template_id"])))
            if t and t.source_type == "template" and t.deleted_at is None:
                return _build_template_response(t)
        except Exception:
            pass

    # Fallback: search for the template created from this source document.
    # Handles workers running old code that didn't write template_id back.
    stmt = (
        select(Document)
        .where(
            Document.source_type == "template",
            Document.deleted_at.is_(None),
            Document.source_metadata["source_document_id"].astext == str(template_id),
        )
        .order_by(Document.created_at.desc())
        .limit(1)
    )
    result = await db.execute(stmt)
    found = result.scalar_one_or_none()
    if found:
        return _build_template_response(found)

    # Not created yet — return the source doc so frontend sees current status
    return _build_template_response(doc)


@router.patch("/{template_id}", response_model=TemplateResponse)
async def update_template(
    template_id: uuid.UUID,
    body: TemplateUpdateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Update template description."""
    doc = await db.get(Document, template_id)
    if not doc or doc.source_type != "template" or doc.deleted_at is not None:
        from fastapi import HTTPException
        raise HTTPException(404, "Template not found")

    if body.description is not None:
        doc.description = body.description

    meta_changed = False
    if body.language_mode is not None or body.sections is not None:
        meta = dict(doc.source_metadata or {})
        if body.language_mode is not None:
            meta["language_mode"] = body.language_mode
        if body.sections is not None:
            meta["sections"] = [s.model_dump() for s in body.sections]
        doc.source_metadata = meta
        meta_changed = True

    if meta_changed:
        from sqlalchemy.orm.attributes import flag_modified
        flag_modified(doc, "source_metadata")

    await db.commit()
    await db.refresh(doc)
    return _build_template_response(doc)


@router.put("/{template_id}/fields", response_model=TemplateResponse)
async def update_template_fields(
    template_id: uuid.UUID,
    body: TemplateFieldsUpdateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Update template field names/placeholders and regenerate DOCX."""
    from docx import Document as DocxDocument
    from src.models.storage import StorageConfig
    from src.services.storage import get_storage_backend

    doc = await db.get(Document, template_id)
    if not doc or doc.source_type != "template" or doc.deleted_at is not None:
        from fastapi import HTTPException
        raise HTTPException(404, "Template not found")

    meta = doc.source_metadata or {}
    old_fields = meta.get("template_fields", [])

    # Build rename map: old_placeholder -> new_placeholder
    rename_map = {}
    new_fields = []
    for update_field in body.fields:
        # Find matching old field
        old = next((f for f in old_fields if f["id"] == update_field.id), None)
        if old and old["placeholder"] != update_field.placeholder:
            rename_map[old["placeholder"]] = update_field.placeholder
        new_type = (update_field.type if update_field.type
                    else (old["type"] if old else "blank"))
        # `select` type requires options — fall back to legacy/no-op if absent
        new_options = (
            update_field.options
            if update_field.options is not None
            else (old.get("options") if old else None)
        )
        new_section = (
            update_field.section_key
            if update_field.section_key is not None
            else (old.get("section_key") if old else None)
        )
        new_required = (
            update_field.required
            if update_field.required is not None
            else (old.get("required") if old else None)
        )
        new_fields.append({
            "id": update_field.id,
            "placeholder": update_field.placeholder,
            "label": update_field.label,
            "description": (update_field.description if update_field.description is not None
                            else (old.get("description", "") if old else "")),
            "location": old["location"] if old else "",
            "type": new_type,
            "options": new_options,
            "section_key": new_section,
            "required": new_required,
        })

    # If any renames, update the DOCX file
    if rename_map:
        storage_cfg = await db.get(StorageConfig, doc.storage_config_id)
        backend = get_storage_backend(storage_cfg)
        doc_bytes = await backend.read(doc.file_path)

        docx = DocxDocument(BytesIO(doc_bytes))

        # Simple text replacement across all paragraphs
        for para in _iter_all_paragraphs(docx):
            text = para.text
            changed = False
            for old_name, new_name in rename_map.items():
                old_token = "{" + old_name + "}"
                new_token = "{" + new_name + "}"
                if old_token in text:
                    text = text.replace(old_token, new_token)
                    changed = True
            if changed and para.runs:
                # Merge runs and set new text
                para.runs[0].text = text
                for r in para.runs[1:]:
                    r.text = ""

        # Save updated DOCX
        output = BytesIO()
        docx.save(output)
        save_result = await backend.save(output.getvalue(), doc.original_filename)
        doc.file_path = save_result.file_path
        doc.file_name = save_result.file_name
        doc.file_size = save_result.file_size
        doc.checksum = save_result.checksum

    # Update metadata (create new dict to trigger SQLAlchemy dirty detection)
    meta = dict(meta)
    meta["template_fields"] = new_fields
    doc.source_metadata = meta

    from sqlalchemy.orm.attributes import flag_modified
    flag_modified(doc, "source_metadata")

    await db.commit()
    await db.refresh(doc)
    return _build_template_response(doc)


@router.put("/{template_id}/file", response_model=TemplateResponse)
async def upload_template_file(
    template_id: uuid.UUID,
    current_user: CurrentUser,
    request: Request,
    db: AsyncSession = Depends(get_db),
):
    """Upload an edited template DOCX file. Re-scans for {placeholder} patterns."""
    from docx import Document as DocxDocument
    from src.models.storage import StorageConfig
    from src.services.storage import get_storage_backend

    doc = await db.get(Document, template_id)
    if not doc or doc.source_type != "template" or doc.deleted_at is not None:
        from fastapi import HTTPException
        raise HTTPException(404, "Template not found")

    # Read raw body (binary DOCX)
    body = await request.body()
    if not body:
        from fastapi import HTTPException
        raise HTTPException(400, "Empty file")

    # Save new file to storage, then clean up the previous version
    storage_cfg = await db.get(StorageConfig, doc.storage_config_id)
    backend = get_storage_backend(storage_cfg)
    old_file_path = doc.file_path
    save_result = await backend.save(body, doc.original_filename)

    doc.file_path = save_result.file_path
    doc.file_name = save_result.file_name
    doc.file_size = save_result.file_size
    doc.checksum = save_result.checksum

    # Best-effort: delete the now-orphaned previous file. Don't fail the request
    # if cleanup errors (the new file is already saved + pointer updated).
    if old_file_path and old_file_path != save_result.file_path:
        try:
            await backend.delete(old_file_path)
        except Exception:
            import logging
            logging.getLogger(__name__).warning(
                "Failed to delete previous template file: %s", old_file_path,
            )

    # Re-scan for {placeholder} patterns in the updated file
    docx = DocxDocument(BytesIO(body))
    placeholder_re = re.compile(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}")

    # Preserve existing label + description across re-upload
    old_meta = doc.source_metadata or {}
    old_fields_map = {
        f["placeholder"]: f for f in old_meta.get("template_fields", [])
    }

    new_fields = []
    field_idx = 0
    for para in _iter_all_paragraphs(docx):
        text = para.text
        for m in placeholder_re.finditer(text):
            name = m.group(1)
            old = old_fields_map.get(name, {})
            new_fields.append({
                "id": f"f{field_idx}",
                "placeholder": name,
                "label": old.get("label") or name.replace("_", " ").title(),
                "description": old.get("description", ""),
                "location": old.get("location", ""),
                "type": old.get("type", "placeholder"),
                "options": old.get("options"),
                "section_key": old.get("section_key"),
                "required": old.get("required"),
            })
            field_idx += 1

    meta = dict(doc.source_metadata or {})
    meta["template_fields"] = new_fields
    meta["extraction_status"] = "completed"
    doc.source_metadata = meta

    from sqlalchemy.orm.attributes import flag_modified
    flag_modified(doc, "source_metadata")

    await db.commit()
    await db.refresh(doc)
    return _build_template_response(doc)


@router.post("/{template_id}/rescan", response_model=TemplateResponse)
async def rescan_template_fields(
    template_id: uuid.UUID,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Re-scan the template DOCX file for {placeholder} patterns and update field list."""
    from docx import Document as DocxDocument
    from src.models.storage import StorageConfig
    from src.services.storage import get_storage_backend

    doc = await db.get(Document, template_id)
    if not doc or doc.source_type != "template" or doc.deleted_at is not None:
        from fastapi import HTTPException
        raise HTTPException(404, "Template not found")

    storage_cfg = await db.get(StorageConfig, doc.storage_config_id)
    backend = get_storage_backend(storage_cfg)
    doc_bytes = await backend.read(doc.file_path)

    docx = DocxDocument(BytesIO(doc_bytes))
    placeholder_re = re.compile(r"\{([a-zA-Z_][a-zA-Z0-9_]*)\}")

    # Build lookup from old fields to preserve label + description
    old_meta = doc.source_metadata or {}
    old_fields = old_meta.get("template_fields", [])
    old_fields_map = {f["placeholder"]: f for f in old_fields}

    new_fields = []
    field_idx = 0
    seen = set()

    for para in _iter_all_paragraphs(docx):
        text = para.text
        for m in placeholder_re.finditer(text):
            name = m.group(1)
            if name not in seen:
                seen.add(name)
                old = old_fields_map.get(name, {})
                new_fields.append({
                    "id": f"f{field_idx}",
                    "placeholder": name,
                    "label": old.get("label") or name.replace("_", " ").title(),
                    "description": old.get("description", ""),
                    "location": old.get("location", ""),
                    "type": old.get("type", "placeholder"),
                    "options": old.get("options"),
                    "section_key": old.get("section_key"),
                    "required": old.get("required"),
                })
                field_idx += 1

    meta = dict(doc.source_metadata or {})
    meta["template_fields"] = new_fields
    doc.source_metadata = meta

    from sqlalchemy.orm.attributes import flag_modified
    flag_modified(doc, "source_metadata")

    await db.commit()
    await db.refresh(doc)
    return _build_template_response(doc)


@router.post("/{document_id}/extract")
async def extract_template_from_document(
    document_id: uuid.UUID,
    current_user: CurrentUser,
    request: Request,
    db: AsyncSession = Depends(get_db),
    description: str | None = Query(None),
):
    """Manually trigger template extraction from an existing document."""
    from fastapi import HTTPException
    from src.api.v1.routes.documents import get_arq_pool
    from src.worker.dispatch import dispatch_task
    from sqlalchemy.orm.attributes import flag_modified

    doc = await db.get(Document, document_id)
    if not doc or doc.deleted_at is not None:
        raise HTTPException(404, "Document not found")

    # Mark as pending immediately so GET /templates/{id} doesn't 404 during polling
    meta = dict(doc.source_metadata or {})
    meta["extraction_status"] = "pending"
    doc.source_metadata = meta
    flag_modified(doc, "source_metadata")
    await db.commit()

    arq_pool = get_arq_pool(request)
    task = await dispatch_task(
        arq_pool=arq_pool,
        func_name="extract_template_task",
        task_name="extract_template",
        db=db,
        owner_id=current_user.id,
        related_type="document",
        related_id=document_id,
        document_id=document_id,
        user_id=current_user.id,
        description=description,
    )
    return {"status": "queued", "task_id": str(task.id)}


# ── Draft/commit flow (Phase B) ─────────────────────────────────────

@router.post("/{document_id}/extract-draft")
async def extract_template_draft_endpoint(
    document_id: uuid.UUID,
    current_user: CurrentUser,
    request: Request,
    db: AsyncSession = Depends(get_db),
):
    """Trigger LLM draft extraction (no template Document created).

    Draft result is stored on BackgroundTask.result and fetched via
    GET /templates/{document_id}/draft. User edits + commits via
    POST /templates/{document_id}/commit.
    """
    from fastapi import HTTPException
    from src.api.v1.routes.documents import get_arq_pool
    from src.worker.dispatch import dispatch_task
    from sqlalchemy.orm.attributes import flag_modified

    doc = await db.get(Document, document_id)
    if not doc or doc.deleted_at is not None:
        raise HTTPException(404, "Document not found")

    meta = dict(doc.source_metadata or {})
    meta["extraction_status"] = "pending"
    meta.pop("extraction_error", None)
    doc.source_metadata = meta
    flag_modified(doc, "source_metadata")
    await db.commit()

    arq_pool = get_arq_pool(request)
    task = await dispatch_task(
        arq_pool=arq_pool,
        func_name="extract_template_draft_task",
        task_name="extract_template_draft",
        db=db,
        owner_id=current_user.id,
        related_type="document",
        related_id=document_id,
        document_id=document_id,
    )
    return {"status": "queued", "task_id": str(task.id)}


@router.get("/{document_id}/draft")
async def get_template_draft(
    document_id: uuid.UUID,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Fetch the most recent draft extraction result for a source document.

    Returns:
      {
        status: "pending" | "running" | "success" | "failure" | "none",
        task_id: "<uuid>" | null,
        draft: { fields, doc_description, source_document_id } | null,
        error: "<message>" | null,
      }
    """
    from fastapi import HTTPException
    from sqlalchemy import select
    from src.models.processing import BackgroundTask

    doc = await db.get(Document, document_id)
    if not doc or doc.deleted_at is not None:
        raise HTTPException(404, "Document not found")
    if doc.owner_id != current_user.id:
        raise HTTPException(403, "Not allowed")

    stmt = (
        select(BackgroundTask)
        .where(
            BackgroundTask.related_type == "document",
            BackgroundTask.related_id == document_id,
            BackgroundTask.task_name == "extract_template_draft",
        )
        .order_by(BackgroundTask.created_at.desc())
        .limit(1)
    )
    result = await db.execute(stmt)
    task = result.scalar_one_or_none()

    if not task:
        return {
            "status": "none",
            "task_id": None,
            "draft": None,
            "error": None,
        }

    draft_payload = None
    err = None
    if task.status == "success" and isinstance(task.result, dict) and not task.result.get("error"):
        draft_payload = task.result
    elif task.status == "failure":
        err = task.error_message or (task.result.get("error") if isinstance(task.result, dict) else None)

    return {
        "status": task.status,
        "task_id": str(task.id),
        "draft": draft_payload,
        "error": err,
    }


@router.post("/{document_id}/commit", response_model=TemplateResponse)
async def commit_template_endpoint(
    document_id: uuid.UUID,
    current_user: CurrentUser,
    body: dict,
    db: AsyncSession = Depends(get_db),
):
    """Commit user-edited draft fields into a template Document.

    Request body:
      {
        "fields": [
          {id, name, label, description, type, current, location},
          ...
        ],
        "doc_description": "optional 1-2 câu mô tả"
      }
    """
    from fastapi import HTTPException
    from src.services.template_service import commit_template

    doc = await db.get(Document, document_id)
    if not doc or doc.deleted_at is not None:
        raise HTTPException(404, "Source document not found")
    if doc.owner_id != current_user.id:
        raise HTTPException(403, "Not allowed")

    fields = body.get("fields") or []
    doc_description = (body.get("doc_description") or "").strip() or None

    if not isinstance(fields, list) or not fields:
        raise HTTPException(400, "Field list is required")

    result = await commit_template(
        db=db,
        source_document_id=document_id,
        fields=fields,
        user_id=current_user.id,
        doc_description=doc_description,
    )
    if result.get("error"):
        raise HTTPException(400, result["error"])

    template_id = uuid.UUID(result["template_id"])
    template_doc = await db.get(Document, template_id)
    return _build_template_response(template_doc)


def _iter_all_paragraphs(docx):
    """Iterate all paragraphs in a DOCX (body + tables + headers/footers)."""
    for para in docx.paragraphs:
        yield para
    for table in docx.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para
    for section in docx.sections:
        for part in [section.header, section.footer]:
            if part:
                for para in part.paragraphs:
                    yield para
