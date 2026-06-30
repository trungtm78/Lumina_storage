"""Document Generator: fill template {placeholders} with field values, save, generate PDF preview."""

from __future__ import annotations

import csv
import io
import json
import uuid
import zipfile
from io import BytesIO
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from src.api.deps import CurrentUser
from src.core.config import get_settings
from src.core.database import get_db
from src.core.template_presets import TEMPLATE_FIELD_PRESETS
from src.models.document import Document
from src.services.reference_loader import (
    format_reference_block,
    load_reference_excerpts,
)
from src.services.generator_docx import (
    _apply_html_edits_to_docx,
    _iter_all_paragraphs,
    _merge_runs,
    _replace_in_runs,
    _validate_field_values,
)
from src.services.storage import get_storage_backend

router = APIRouter(prefix="/generator", tags=["generator"])


class GenerateRequest(BaseModel):
    template_id: str
    field_values: dict[str, str]
    output_filename: str | None = None
    folder_id: str | None = None


class GenerateResponse(BaseModel):
    document_id: str
    original_filename: str
    preview_pdf_id: str | None
    applied_count: int



@router.post("/generate", response_model=GenerateResponse)
async def generate_document(
    body: GenerateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    from src.models.storage import StorageConfig

    svc = GeneratorService(db)
    rendered_doc, out_filename, applied_count = await svc.execute_generate(
        template_id=uuid.UUID(body.template_id),
        field_values=body.field_values,
        output_filename=body.output_filename,
        folder_id=uuid.UUID(body.folder_id) if body.folder_id else None,
        owner_id=current_user.id,
    )
    # Phase 3 — COMMIT CỐ Ý: document phải BỀN trước khối preview best-effort bên dưới.
    # Preview add/flush pdf_doc trong `try/except: pass`; nếu flush lỗi DB → transaction
    # bị poison nhưng exception bị nuốt → boundary cuối request sẽ rollback và MẤT luôn
    # document user vừa tạo. Commit ở đây tách document khỏi side-effect preview. KHÔNG gỡ.
    await db.commit()

    document_id = rendered_doc.id

    # Generate PDF preview via Gotenberg (best-effort)
    preview_pdf_id = None
    storage_cfg = await db.get(StorageConfig, rendered_doc.storage_config_id)
    settings = get_settings()
    if settings.gotenberg_url and storage_cfg:
        try:
            backend = get_storage_backend(storage_cfg)
            rendered_bytes = await backend.read(rendered_doc.file_path)
            src_path = Path(out_filename)
            import httpx
            async with httpx.AsyncClient(timeout=60) as client:
                resp = await client.post(
                    f"{settings.gotenberg_url}/forms/libreoffice/convert",
                    files={"files": (out_filename, rendered_bytes, rendered_doc.mime_type)},
                )
                if resp.status_code == 200:
                    pdf_bytes = resp.content
                    pdf_filename = f"{src_path.stem}_generated_preview.pdf"
                    pdf_result = await backend.save(pdf_bytes, pdf_filename)

                    pdf_doc = Document(
                        title=f"{rendered_doc.title}_preview",
                        file_name=pdf_result.file_name,
                        original_filename=pdf_filename,
                        file_path=pdf_result.file_path,
                        file_size=pdf_result.file_size,
                        mime_type="application/pdf",
                        extension="pdf",
                        checksum=pdf_result.checksum,
                        storage_config_id=rendered_doc.storage_config_id,
                        owner_id=current_user.id,
                        source_type="skill_temp",
                        source_metadata={"source_document_id": str(document_id)},
                    )
                    db.add(pdf_doc)
                    await db.flush()
                    await db.refresh(pdf_doc)
                    # Phase 3 — COMMIT CỐ Ý: persist preview NGAY (best-effort, độc lập
                    # document đã commit ở trên). Nếu bước nào của preview lỗi → except
                    # nuốt, document vẫn bền. KHÔNG gỡ.
                    await db.commit()
                    preview_pdf_id = pdf_doc.id
        except Exception:
            pass

    return GenerateResponse(
        document_id=str(document_id),
        original_filename=out_filename,
        preview_pdf_id=str(preview_pdf_id) if preview_pdf_id else None,
        applied_count=applied_count,
    )


class RenderPdfRequest(BaseModel):
    template_id: str
    field_values: dict[str, str]


@router.post("/render-pdf")
async def render_pdf_preview(
    body: RenderPdfRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Fill template placeholders and return a PDF preview — nothing is saved."""
    from docx import Document as DocxDocument
    from src.models.storage import StorageConfig

    template_doc = await db.get(Document, uuid.UUID(body.template_id))
    if not template_doc or template_doc.source_type != "template" or template_doc.deleted_at is not None:
        raise HTTPException(404, "Template not found")

    storage_cfg = await db.get(StorageConfig, template_doc.storage_config_id)
    if not storage_cfg:
        raise HTTPException(500, "Storage config not found")

    backend = get_storage_backend(storage_cfg)
    doc_bytes = await backend.read(template_doc.file_path)

    docx = DocxDocument(BytesIO(doc_bytes))
    for para in _iter_all_paragraphs(docx):
        _merge_runs(para)
        if not para.runs:
            continue
        text = para.runs[0].text
        changed = False
        for key, value in body.field_values.items():
            token = "{" + key + "}"
            if token in text:
                text = text.replace(token, value)
                changed = True
        if changed:
            para.runs[0].text = text

    rendered = BytesIO()
    docx.save(rendered)
    rendered_bytes = rendered.getvalue()

    src_path = Path(template_doc.original_filename)
    settings = get_settings()
    if not settings.gotenberg_url:
        raise HTTPException(503, "PDF preview unavailable — Gotenberg not configured")

    try:
        import httpx
        async with httpx.AsyncClient(timeout=60) as client:
            resp = await client.post(
                f"{settings.gotenberg_url}/forms/libreoffice/convert",
                files={"files": (src_path.name, rendered_bytes, template_doc.mime_type)},
            )
        if resp.status_code != 200:
            raise HTTPException(502, "PDF conversion failed")
        return StreamingResponse(
            BytesIO(resp.content),
            media_type="application/pdf",
            headers={"Content-Disposition": f'inline; filename="{src_path.stem}_preview.pdf"'},
        )
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(502, "PDF conversion failed")


@router.get("/field-presets")
async def get_field_presets(_: CurrentUser):
    """Return curated dropdown values for `select` template fields.

    Frontend uses these as suggestions when configuring a `select` field.
    """
    return TEMPLATE_FIELD_PRESETS


class DraftRequest(BaseModel):
    doc_type: str
    description: str = ""
    # Phase 3 — multi-select danh sách phụ lục/tài liệu đã duyệt làm reference
    # cho LLM. Chỉ định `language` để buộc đầu ra đơn ngữ hoặc song ngữ.
    reference_document_ids: list[str] = []
    language: str | None = None  # "single" | "bilingual" | None


class DraftResponse(BaseModel):
    content: str
    version: int
    label: str
    # IDs of reference docs whose text actually made it into the LLM prompt.
    # Subset of the requested `reference_document_ids` — missing/empty/forbidden
    # docs are silently skipped so the draft still succeeds.
    references_used: list[str] = []


class ReviseRequest(BaseModel):
    content: str
    instruction: str
    version: int


class ReviseResponse(BaseModel):
    content: str
    version: int
    label: str


_DOC_TYPE_LABELS: dict[str, str] = {
    "sales_contract": "hợp đồng bán hàng",
    "purchase_contract": "hợp đồng mua hàng",
    "nda": "thỏa thuận bảo mật thông tin (NDA)",
    "internal_memo": "công văn nội bộ",
    "proposal": "đề xuất hợp tác",
    "report": "báo cáo",
}


def _language_directive(language: str | None) -> str:
    if language == "bilingual":
        return (
            "\n\nVăn bản phải SONG NGỮ Việt-Anh: mỗi điều khoản viết tiếng Việt "
            "trước, ngay dưới là bản dịch tiếng Anh tương ứng."
        )
    if language == "single":
        return "\n\nVăn bản chỉ dùng tiếng Việt."
    return ""


@router.post("/draft", response_model=DraftResponse)
async def draft_document(
    body: DraftRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    settings = get_settings()
    doc_label = _DOC_TYPE_LABELS.get(body.doc_type, body.doc_type)

    excerpts = await load_reference_excerpts(
        db, body.reference_document_ids, current_user.id,
    )
    reference_block = format_reference_block(excerpts)
    language_directive = _language_directive(body.language)

    user_prompt = f"Soạn thảo {doc_label}."
    if body.description:
        user_prompt += f"\n\nYêu cầu bổ sung: {body.description}"
    if reference_block:
        user_prompt += f"\n\n{reference_block}"
    user_prompt += language_directive
    user_prompt += (
        "\n\nSử dụng các placeholder như {{ten_cong_ty}}, {{ngay_ky}}, "
        "{{so_tien}}, v.v. ở những chỗ phù hợp."
    )

    messages = [
        {
            "role": "system",
            "content": (
                "Bạn là chuyên gia soạn thảo văn bản kinh doanh cho Got It Vietnam — "
                "công ty cung cấp giải pháp quà tặng số B2B. "
                "Hãy soạn văn bản chuyên nghiệp bằng tiếng Việt với các token placeholder dạng {tên_placeholder} "
                "cho những thông tin biến động như tên công ty, ngày tháng, số tiền, v.v. "
                "Trả về nội dung văn bản thuần túy, không có giải thích thêm."
            ),
        },
        {"role": "user", "content": user_prompt},
    ]

    from src.ai import AIGateway
    response = await AIGateway(db).complete(messages)
    import re as _re
    content: str = response.choices[0].message.content or ""
    content = _re.sub(r"\{\{(\w+)\}\}", r"{\1}", content)

    return DraftResponse(
        content=content,
        version=1,
        label="Phiên bản 1",
        references_used=[str(e.document_id) for e in excerpts],
    )


@router.post("/revise", response_model=ReviseResponse)
async def revise_document(
    body: ReviseRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    from src.ai import AIGateway

    messages = [
        {
            "role": "system",
            "content": (
                "Bạn là chuyên gia chỉnh sửa văn bản kinh doanh tiếng Việt. "
                "Hãy chỉnh sửa văn bản theo yêu cầu, giữ nguyên định dạng và các token placeholder dạng {tên_placeholder}. "
                "Trả về nội dung văn bản đã chỉnh sửa, không có giải thích thêm."
            ),
        },
        {
            "role": "user",
            "content": (
                f"Văn bản hiện tại:\n\n{body.content}\n\n"
                f"Yêu cầu chỉnh sửa: {body.instruction}\n\n"
                "Hãy trả về văn bản đã được chỉnh sửa."
            ),
        },
    ]

    response = await AIGateway(db).complete(messages)
    import re as _re
    content: str = response.choices[0].message.content or ""
    content = _re.sub(r"\{\{(\w+)\}\}", r"{\1}", content)

    next_ver = body.version + 1
    label = "Phiên bản cuối" if next_ver >= 3 else f"Phiên bản {next_ver}"

    return ReviseResponse(content=content, version=next_ver, label=label)


def _parse_uploaded_file(content: bytes, filename: str) -> tuple[list[str], list[dict[str, str]]]:
    """Return (column_names, rows) from xlsx or csv content."""
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = [dict(r) for r in reader]
        columns = list(reader.fieldnames or (rows[0].keys() if rows else []))
        return columns, rows
    # xlsx
    from openpyxl import load_workbook
    wb = load_workbook(filename=BytesIO(content), read_only=True)
    ws = wb.active
    headers: list[str] = []
    rows: list[dict[str, str]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c or "") for c in row]
        else:
            rows.append({headers[j]: str(v or "") for j, v in enumerate(row)})
    wb.close()
    return headers, rows


class MapColumnsResponse(BaseModel):
    columns: list[str]
    mapping: list[dict[str, str]]
    sample_row: dict[str, str]
    total_rows: int


@router.post("/map-columns", response_model=MapColumnsResponse)
async def map_columns(
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
    template_id: str = Form(...),
    instruction: str = Form(""),
    file: UploadFile = File(...),
):
    template_doc = await db.get(Document, uuid.UUID(template_id))
    if not template_doc or template_doc.source_type != "template" or template_doc.deleted_at is not None:
        raise HTTPException(404, "Template not found")

    meta = template_doc.source_metadata or {}
    placeholders = [f["placeholder"] for f in meta.get("template_fields", [])]

    content = await file.read()
    columns, rows = _parse_uploaded_file(content, file.filename or "upload.xlsx")

    if not rows:
        raise HTTPException(400, "File is empty or has no data rows")

    sample_row = rows[0]

    # Use AI to map columns → placeholders
    settings = get_settings()
    prompt = (
        f"Bạn là chuyên gia ánh xạ dữ liệu.\n"
        f"Các cột trong file: {json.dumps(columns, ensure_ascii=False)}\n"
        f"Các placeholder trong tài liệu: {json.dumps(placeholders, ensure_ascii=False)}\n"
        + (f"Hướng dẫn bổ sung: {instruction}\n" if instruction else "")
        + "Hãy trả về ĐÚNG một JSON array (không có giải thích) ánh xạ từng cột sang placeholder phù hợp nhất.\n"
        "Ví dụ: [{\"column\": \"Tên KH\", \"placeholder\": \"ten_khach_hang\"}, ...]\n"
        "Chỉ ánh xạ các cột có placeholder tương ứng. Bỏ qua cột không có placeholder phù hợp."
    )

    from src.ai import AIGateway
    response = await AIGateway(db).complete([
        {"role": "system", "content": "Bạn chỉ trả về JSON, không có markdown hay giải thích."},
        {"role": "user", "content": prompt},
    ])
    raw = (response.choices[0].message.content or "").strip()
    # strip markdown fences if any
    if raw.startswith("```"):
        raw = "\n".join(raw.split("\n")[1:])
        raw = raw.rsplit("```", 1)[0].strip()

    try:
        mapping: list[dict[str, str]] = json.loads(raw)
    except Exception:
        mapping = []

    return MapColumnsResponse(
        columns=columns,
        mapping=mapping,
        sample_row={m["column"]: sample_row.get(m["column"], "") for m in mapping},
        total_rows=len(rows),
    )


@router.post("/batch")
async def batch_generate_documents(
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
    template_id: str = Form(...),
    column_mapping: str = Form(...),  # JSON string
    file: UploadFile = File(...),
):
    from docx import Document as DocxDocument
    from src.models.storage import StorageConfig

    template_doc = await db.get(Document, uuid.UUID(template_id))
    if not template_doc or template_doc.source_type != "template" or template_doc.deleted_at is not None:
        raise HTTPException(404, "Template not found")

    storage_cfg = await db.get(StorageConfig, template_doc.storage_config_id)
    if not storage_cfg:
        raise HTTPException(500, "Storage config not found")

    backend = get_storage_backend(storage_cfg)
    template_bytes = await backend.read(template_doc.file_path)

    try:
        mapping: list[dict[str, str]] = json.loads(column_mapping)
    except Exception:
        raise HTTPException(400, "Invalid column_mapping JSON")

    # col → placeholder lookup
    col_to_placeholder = {m["column"]: m["placeholder"] for m in mapping}

    content = await file.read()
    _columns, rows = _parse_uploaded_file(content, file.filename or "upload.xlsx")

    if not rows:
        raise HTTPException(400, "File has no data rows")

    src_path = Path(template_doc.original_filename)
    zip_buffer = BytesIO()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for idx, row in enumerate(rows, start=1):
            field_values = {
                placeholder: row.get(col, "")
                for col, placeholder in col_to_placeholder.items()
            }
            docx = DocxDocument(BytesIO(template_bytes))
            for para in _iter_all_paragraphs(docx):
                _merge_runs(para)
                if not para.runs:
                    continue
                text = para.runs[0].text
                changed = False
                for key, value in field_values.items():
                    token = "{" + key + "}"
                    if token in text:
                        text = text.replace(token, value)
                        changed = True
                if changed:
                    para.runs[0].text = text

            out = BytesIO()
            docx.save(out)
            zf.writestr(f"{src_path.stem}_{idx:03d}{src_path.suffix}", out.getvalue())

    zip_buffer.seek(0)
    filename = f"{src_path.stem}_batch_{len(rows)}_documents.zip"
    from urllib.parse import quote
    encoded = quote(filename, safe="")
    ascii_fallback = filename.encode("ascii", errors="replace").decode("ascii")
    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{ascii_fallback}"; '
                f"filename*=UTF-8''{encoded}"
            )
        },
    )


class DraftToTemplateRequest(BaseModel):
    content: str
    title: str
    description: str = ""


class DraftToTemplateResponse(BaseModel):
    template_id: str
    title: str
    field_count: int
    template_fields: list[dict]


@router.post("/draft-to-template", response_model=DraftToTemplateResponse)
async def draft_to_template(
    body: DraftToTemplateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Convert an AI-drafted text (with {placeholder} tokens) into a saved DOCX template."""
    import re
    from docx import Document as DocxDocument
    from src.models.storage import StorageConfig
    from sqlalchemy import select as sa_select

    # Normalize {{placeholder}} → {placeholder} (AI sometimes generates double braces)
    normalized = re.sub(r"\{\{(\w+)\}\}", r"{\1}", body.content)

    # Build DOCX from text content
    docx = DocxDocument()
    for line in normalized.split("\n"):
        docx.add_paragraph(line)
    buf = BytesIO()
    docx.save(buf)
    docx_bytes = buf.getvalue()

    # Extract {placeholder} tokens
    tokens = re.findall(r"\{([^{}]+)\}", normalized)
    unique_tokens = list(dict.fromkeys(tokens))

    template_fields = [
        {
            "placeholder": p,
            "label": p.replace("_", " ").title(),
            "description": "",
            "type": "date" if any(x in p for x in ("date", "ngay", "ngày", "time")) else "placeholder",
            "location": "body",
        }
        for p in unique_tokens
    ]

    # Find default storage config
    stmt = sa_select(StorageConfig).where(StorageConfig.is_default.is_(True)).limit(1)
    result = await db.execute(stmt)
    storage_cfg = result.scalar_one_or_none()
    if not storage_cfg:
        stmt = sa_select(StorageConfig).limit(1)
        result = await db.execute(stmt)
        storage_cfg = result.scalar_one_or_none()
    if not storage_cfg:
        raise HTTPException(500, "No storage configuration available")

    backend = get_storage_backend(storage_cfg)
    filename = f"{body.title}.docx"
    save_result = await backend.save(docx_bytes, filename)

    template_doc = Document(
        title=body.title,
        file_name=save_result.file_name,
        original_filename=filename,
        file_path=save_result.file_path,
        file_size=save_result.file_size,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        extension="docx",
        checksum=save_result.checksum,
        storage_config_id=storage_cfg.id,
        owner_id=current_user.id,
        source_type="template",
        source_metadata={
            "description": body.description,
            "template_fields": template_fields,
            "extraction_status": "completed",
        },
    )
    db.add(template_doc)
    await db.flush()
    await db.refresh(template_doc)
    # Phase 3: commit ở boundary (get_db).

    return DraftToTemplateResponse(
        template_id=str(template_doc.id),
        title=body.title,
        field_count=len(unique_tokens),
        template_fields=template_fields,
    )


class ExtractFromFileResponse(BaseModel):
    field_values: dict[str, str]
    extracted_count: int
    # Per-field source attribution: placeholder → filename that provided the value
    sources: dict[str, str] = {}
    # Files that couldn't be converted to text (empty or unsupported)
    skipped_files: list[str] = []


@router.post("/extract-from-file", response_model=ExtractFromFileResponse)
async def extract_from_file(
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
    template_id: str = Form(...),
    files: list[UploadFile] = File(...),
):
    """Extract template field values by aggregating content from one or more source files.

    Accepts multiple files (brief, email, quote, scan, etc.). LLM reads all of them
    together and resolves conflicts when the same info appears in multiple places.
    """
    template_doc = await db.get(Document, uuid.UUID(template_id))
    if not template_doc or template_doc.source_type != "template" or template_doc.deleted_at is not None:
        raise HTTPException(404, "Template not found")

    meta = template_doc.source_metadata or {}
    template_fields = meta.get("template_fields", [])
    if not template_fields:
        return ExtractFromFileResponse(field_values={}, extracted_count=0)

    if not files:
        raise HTTPException(400, "At least one file is required")

    # Step 1: convert every file to plain text via MarkItDown
    from markitdown import MarkItDown
    md = MarkItDown()

    sources: list[dict] = []  # [{name, text}]
    skipped: list[str] = []
    for idx, f in enumerate(files):
        content = await f.read()
        filename = f.filename or f"file_{idx + 1}"
        ext = Path(filename).suffix.lower()
        try:
            result = md.convert_stream(BytesIO(content), file_extension=ext)
            text = (result.text_content or "").strip()
        except Exception:
            text = content.decode("utf-8", errors="ignore").strip()

        if not text:
            skipped.append(filename)
            continue
        sources.append({"name": filename, "text": text})

    if not sources:
        return ExtractFromFileResponse(
            field_values={},
            extracted_count=0,
            sources={},
            skipped_files=skipped,
        )

    # Step 2: build combined source block with clear boundaries
    source_block_parts = []
    for i, src in enumerate(sources, start=1):
        source_block_parts.append(
            f"--- FILE {i}: {src['name']} ---\n{src['text']}"
        )
    combined_sources = "\n\n".join(source_block_parts)
    source_filenames = [s["name"] for s in sources]

    valid_placeholders = {f["placeholder"] for f in template_fields}
    field_hints = json.dumps(
        [
            {
                "placeholder": f["placeholder"],
                "label": f.get("label", f["placeholder"]),
                "description": f.get("description", ""),
            }
            for f in template_fields
        ],
        ensure_ascii=False,
        indent=2,
    )

    settings = get_settings()
    multi_hint = (
        "Hãy TỔNG HỢP thông tin từ TẤT CẢ các file trên để điền các trường.\n"
        "- Nếu cùng một thông tin xuất hiện ở nhiều file, ưu tiên giá trị CỤ THỂ hơn + ĐẦY ĐỦ hơn.\n"
        "- Nếu có xung đột rõ rệt (ví dụ 2 tên công ty khác nhau), chọn giá trị xuất hiện ở file trang trọng "
        "hơn (hợp đồng > brief > email).\n"
        "- Với mỗi trường, ghi nhận tên file mà bạn trích xuất thông tin (nếu kết hợp nhiều nguồn thì ghi "
        "file chính).\n"
    ) if len(sources) > 1 else ""

    prompt = (
        "Bạn là chuyên gia trích xuất thông tin từ văn bản.\n\n"
        f"Các trường cần điền (placeholder, nhãn hiển thị, mô tả chi tiết):\n{field_hints}\n\n"
        f"Các file nguồn ({len(sources)} file):\n{combined_sources}\n\n"
        f"{multi_hint}"
        "Hãy trả về ĐÚNG một JSON object có 2 key:\n"
        "  - \"values\": { <placeholder>: <giá_trị_trích_xuất> }\n"
        "  - \"sources\": { <placeholder>: <tên_file> }\n"
        "Chỉ điền những placeholder mà bạn tìm thấy thông tin rõ ràng. "
        "Không điền nếu không chắc chắn hoặc không có thông tin."
    )

    from src.ai import AIGateway
    response = await AIGateway(db).complete([
        {"role": "system", "content": "Bạn chỉ trả về JSON thuần, không có markdown hay giải thích."},
        {"role": "user", "content": prompt},
    ])
    raw = (response.choices[0].message.content or "").strip()
    if raw.startswith("```"):
        raw = "\n".join(raw.split("\n")[1:])
        raw = raw.rsplit("```", 1)[0].strip()

    field_values: dict[str, str] = {}
    field_sources: dict[str, str] = {}
    try:
        extracted: dict = json.loads(raw)
        # New shape: {values: {...}, sources: {...}}
        if "values" in extracted and isinstance(extracted["values"], dict):
            values_raw = extracted["values"]
            sources_raw = extracted.get("sources", {}) if isinstance(extracted.get("sources"), dict) else {}
        else:
            # Fallback: LLM returned old flat shape
            values_raw = extracted
            sources_raw = {}

        for k, v in values_raw.items():
            if k in valid_placeholders and v:
                field_values[k] = str(v)
                src_name = sources_raw.get(k)
                # Trust only filenames we actually sent
                if src_name in source_filenames:
                    field_sources[k] = src_name
                elif len(sources) == 1:
                    # Single-file case: attribute to that file
                    field_sources[k] = source_filenames[0]
    except Exception:
        pass

    return ExtractFromFileResponse(
        field_values=field_values,
        extracted_count=len(field_values),
        sources=field_sources,
        skipped_files=skipped,
    )


class FieldHint(BaseModel):
    placeholder: str
    label: str
    description: str = ""


class ExtractFromTextRequest(BaseModel):
    template_id: str | None = None
    field_hints: list[FieldHint] | None = None
    text: str


@router.post("/extract-from-text", response_model=ExtractFromFileResponse)
async def extract_from_text(
    body: ExtractFromTextRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Extract template field values from pasted plain text (no file upload needed).

    Accepts either a template_id (to load field definitions from DB) or an explicit
    field_hints list (when no template is selected — e.g. fallback FIELDS_BY_TYPE).
    """
    if not body.text.strip():
        raise HTTPException(400, "text must not be empty")

    template_fields: list[dict] = []

    if body.template_id:
        template_doc = await db.get(Document, uuid.UUID(body.template_id))
        if not template_doc or template_doc.source_type != "template" or template_doc.deleted_at is not None:
            raise HTTPException(404, "Template not found")
        meta = template_doc.source_metadata or {}
        template_fields = meta.get("template_fields", [])
    elif body.field_hints:
        template_fields = [
            {"placeholder": h.placeholder if h.placeholder.startswith("{") else "{" + h.placeholder + "}", "label": h.label, "description": h.description}
            for h in body.field_hints
        ]

    if not template_fields:
        return ExtractFromFileResponse(field_values={}, extracted_count=0)

    valid_placeholders = {f["placeholder"] for f in template_fields}
    field_hints_json = json.dumps(
        [
            {
                "placeholder": f["placeholder"],
                "label": f.get("label", f["placeholder"]),
                "description": f.get("description", ""),
            }
            for f in template_fields
        ],
        ensure_ascii=False,
        indent=2,
    )

    prompt = (
        "Bạn là chuyên gia trích xuất thông tin từ văn bản.\n\n"
        f"Các trường cần điền (placeholder, nhãn hiển thị, mô tả chi tiết):\n{field_hints_json}\n\n"
        f"Nội dung nguồn:\n{body.text}\n\n"
        "Hãy trả về ĐÚNG một JSON object có 2 key:\n"
        "  - \"values\": { <placeholder>: <giá_trị_trích_xuất> }\n"
        "  - \"sources\": { <placeholder>: \"text\" }\n"
        "Chỉ điền những placeholder mà bạn tìm thấy thông tin rõ ràng. "
        "Không điền nếu không chắc chắn hoặc không có thông tin."
    )

    from src.ai import AIGateway
    response = await AIGateway(db).complete([
        {"role": "system", "content": "Bạn chỉ trả về JSON thuần, không có markdown hay giải thích."},
        {"role": "user", "content": prompt},
    ])
    raw = (response.choices[0].message.content or "").strip()
    if raw.startswith("```"):
        raw = "\n".join(raw.split("\n")[1:])
        raw = raw.rsplit("```", 1)[0].strip()

    field_values: dict[str, str] = {}
    try:
        extracted: dict = json.loads(raw)
        if "values" in extracted and isinstance(extracted["values"], dict):
            values_raw = extracted["values"]
        else:
            values_raw = extracted

        for k, v in values_raw.items():
            if k in valid_placeholders and v:
                field_values[k] = str(v)
    except Exception:
        pass

    return ExtractFromFileResponse(
        field_values=field_values,
        extracted_count=len(field_values),
        sources={k: "text" for k in field_values},
    )


# ─── Generator Sessions ────────────────────────────────────────────────────────

from src.models.generator import GeneratorSession
from src.services.generator_service import GeneratorService
from src.schemas.generator import (
    AiReviseApplyRequest,
    AiReviseApplyResponse,
    AiReviseRequest,
    AiReviseResponse,
    BlockEditOp,
    DocumentToTemplateRequest,
    DocumentToTemplateResponse,
    GeneratorSessionCreateRequest,
    GeneratorSessionGenerateRequest,
    GeneratorSessionListResponse,
    GeneratorSessionResponse,
    GeneratorSessionUpdateRequest,
    GeneratorSessionVersionCreateRequest,
    GeneratorSessionVersionListResponse,
    GeneratorSessionVersionResponse,
    GeneratorSessionVersionUpdateRequest,
)


from src.services.generator_html import (
    _apply_block_ops,
    _build_diff_html,
    _html_blocks,
    _html_to_docx_bytes,
    _html_to_docx_via_gotenberg,
    _html_to_pdf_bytes,
    _sanitize_html,
    _structural_guard,
    _substitute_fields_html,
    _tokens_in,
    _wrap_html_document,
)


@router.post("/sessions", response_model=GeneratorSessionResponse)
async def create_session(
    body: GeneratorSessionCreateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    svc = GeneratorService(db)
    session = await svc.create_session({
        "user_id": current_user.id,
        "template_id": body.template_id,
        "doc_type": body.doc_type,
        "field_values": body.field_values,
        "title": body.title,
        "folder_id": body.folder_id,
        "status": "draft",
    })
    # Phase 3: commit ở boundary (get_db). repo.create đã flush → id sẵn.
    await db.flush()
    return GeneratorSessionResponse.model_validate(session)


@router.get("/sessions", response_model=GeneratorSessionListResponse)
async def list_sessions(
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
    status: str | None = Query(None),
    limit: int = Query(50, ge=1, le=100),
    offset: int = Query(0, ge=0),
):
    svc = GeneratorService(db)
    items, total = await svc.list_sessions_for_user(
        user_id=current_user.id,
        status=status,
        limit=limit,
        offset=offset,
    )
    # Batch-check which template_ids are still alive (not soft-deleted)
    from sqlalchemy import select as _sa_select
    template_ids = {s.template_id for s in items if s.template_id is not None}
    alive_template_ids: set = set()
    if template_ids:
        from src.models.document import Document as _Doc
        result = await db.execute(
            _sa_select(_Doc.id).where(
                _Doc.id.in_(template_ids),
                _Doc.deleted_at.is_(None),
            )
        )
        alive_template_ids = {row[0] for row in result.all()}

    def _to_response(s) -> GeneratorSessionResponse:
        r = GeneratorSessionResponse.model_validate(s)
        if s.template_id is not None:
            r.template_exists = s.template_id in alive_template_ids
        return r

    return GeneratorSessionListResponse(
        items=[_to_response(s) for s in items],
        total=total,
    )


@router.patch("/sessions/{session_id}", response_model=GeneratorSessionResponse)
async def update_session(
    session_id: uuid.UUID,
    body: GeneratorSessionUpdateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    svc = GeneratorService(db)
    session = await svc.get_session_for_user(session_id, current_user.id)
    if session is None:
        raise HTTPException(404, "Session not found")

    update_data: dict = {}
    if body.field_values is not None:
        update_data["field_values"] = body.field_values
    if body.title is not None:
        update_data["title"] = body.title
    if "folder_id" in body.model_fields_set:
        update_data["folder_id"] = body.folder_id  # None = clear folder
    if body.edited_html is not None:
        update_data["edited_html"] = body.edited_html

    if update_data:
        session = await svc.update_session(session_id, update_data)

    # Phase 3: commit ở boundary (get_db).
    await db.flush()
    return GeneratorSessionResponse.model_validate(session)


@router.post("/sessions/{session_id}/generate", response_model=GeneratorSessionResponse)
async def generate_from_session(
    session_id: uuid.UUID,
    body: GeneratorSessionGenerateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    svc = GeneratorService(db)
    session = await svc.get_session_for_user(session_id, current_user.id)
    if session is None:
        raise HTTPException(404, "Session not found")
    _has_edit = bool(body.version_id or body.edited_html or session.edited_html)
    if session.template_id is None and not _has_edit:
        raise HTTPException(400, "Session has no template_id — cannot generate")

    want_pdf = (body.output_format or "docx").lower() == "pdf"
    # Folder: explicit request wins, else fall back to the folder saved on the draft.
    target_folder_id = body.folder_id or session.folder_id

    # ── Resolve nội dung SỬA TAY (nếu có): version_id > body.edited_html > session.edited_html ──
    edited_html = None
    edited_fv = session.field_values or {}
    if body.version_id:
        ver = await svc.get_version_for_session(body.version_id, session_id)
        if ver is None:
            raise HTTPException(404, "Version not found")
        edited_html = ver.edited_html
        edited_fv = ver.field_values or edited_fv
    elif body.edited_html is not None:
        edited_html = body.edited_html
    elif session.edited_html:
        edited_html = session.edited_html

    # ── NHÁNH SỬA TAY ────────────────────────────────────────────────────────────
    # Khi có edited_html VÀ template_id: patch DOCX gốc (giữ formatting hoàn toàn)
    # Khi chỉ có edited_html (không có template): fallback HTML→DOCX (lossy)
    if edited_html:
        try:
            from sqlalchemy import select as _sa_select
            from src.models.storage import StorageConfig

            _cfg = (
                await db.execute(_sa_select(StorageConfig).where(StorageConfig.is_default.is_(True)).limit(1))
            ).scalar_one_or_none()
            if _cfg is None:
                _cfg = (await db.execute(_sa_select(StorageConfig).limit(1))).scalar_one_or_none()
            if _cfg is None:
                raise HTTPException(500, "No storage configuration available")
            backend = get_storage_backend(_cfg)

            base_name = Path(body.output_filename).stem if body.output_filename else (session.title or "tai_lieu")
            settings = get_settings()

            if session.template_id is not None:
                # ── Đường chính: patch DOCX gốc, giữ toàn bộ formatting ──────────
                # _apply_html_edits_to_docx dùng difflib để tìm paragraph thay đổi,
                # thay text trong DOCX runs (giữ bold/italic/font/size/màu) — không
                # convert HTML→DOCX, không mất style gốc.
                from src.models.document import Document as DocModel
                template_doc = await db.get(DocModel, session.template_id)
                if template_doc is None or template_doc.deleted_at is not None:
                    raise HTTPException(404, "Template not found")
                tmpl_storage_cfg = await db.get(StorageConfig, template_doc.storage_config_id)
                tmpl_backend = get_storage_backend(tmpl_storage_cfg)
                template_bytes = await tmpl_backend.read(template_doc.file_path)

                docx_bytes = _apply_html_edits_to_docx(template_bytes, edited_html, edited_fv)

                if want_pdf:
                    if not settings.gotenberg_url:
                        raise HTTPException(503, "PDF conversion unavailable — Gotenberg not configured")
                    async with __import__("httpx").AsyncClient(timeout=120) as client:
                        resp = await client.post(
                            f"{settings.gotenberg_url}/forms/libreoffice/convert",
                            files={"files": (f"{base_name}.docx", docx_bytes,
                                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                        )
                    if resp.status_code != 200:
                        raise HTTPException(502, "DOCX→PDF conversion failed")
                    file_bytes = resp.content
                    filename, mime, ext = f"{base_name}.pdf", "application/pdf", "pdf"
                else:
                    file_bytes = docx_bytes
                    filename = f"{base_name}.docx"
                    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    ext = "docx"

            else:
                # ── Fallback: không có template, convert HTML→DOCX (lossy) ────────
                substituted = _substitute_fields_html(edited_html, edited_fv)
                if want_pdf:
                    if not settings.gotenberg_url:
                        raise HTTPException(503, "PDF conversion unavailable — Gotenberg not configured")
                    file_bytes = await _html_to_pdf_bytes(_wrap_html_document(substituted), settings.gotenberg_url)
                    filename, mime, ext = f"{base_name}.pdf", "application/pdf", "pdf"
                else:
                    if settings.gotenberg_url:
                        file_bytes = await _html_to_docx_via_gotenberg(
                            _wrap_html_document(substituted), settings.gotenberg_url
                        )
                    else:
                        file_bytes = _html_to_docx_bytes(substituted)
                    filename = f"{base_name}.docx"
                    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    ext = "docx"

            save_result = await backend.save(file_bytes, filename)
            doc = Document(
                title=base_name,
                file_name=save_result.file_name,
                original_filename=filename,
                file_path=save_result.file_path,
                file_size=save_result.file_size,
                mime_type=mime,
                extension=ext,
                checksum=save_result.checksum,
                storage_config_id=_cfg.id,
                owner_id=current_user.id,
                folder_id=target_folder_id,
                source_type="generated",
                source_metadata={
                    "template_id": str(session.template_id) if session.template_id else None,
                    "from_manual_edit": True,
                },
            )
            db.add(doc)
            await db.flush()
            await db.refresh(doc)
            session = await svc.update_session(session_id, {
                "status": "completed",
                "document_id": doc.id,
                "folder_id": target_folder_id,
                "edited_html": edited_html,
            })
            # Phase 3 — COMMIT CỐ Ý (status-machine): persist trạng thái completed/failed
            # của session. Nhánh except ghi status="failed" RỒI commit TRƯỚC `raise` →
            # phải bền qua rollback của boundary. KHÔNG gỡ.
            await db.commit()
        except HTTPException:
            await svc.update_session(session_id, {"status": "failed", "error_message": "Manual-edit generation failed"})
            await db.commit()
            raise
        return GeneratorSessionResponse.model_validate(session)

    try:
        rendered_doc, out_filename, _count = await svc.execute_generate(
            template_id=session.template_id,
            field_values=session.field_values or {},
            output_filename=body.output_filename,
            folder_id=target_folder_id,
            owner_id=current_user.id,
            skip_field_validation=body.skip_field_validation,
        )

        final_document_id = rendered_doc.id

        # Convert to PDF via Gotenberg if requested
        if want_pdf:
            settings = get_settings()
            if not settings.gotenberg_url:
                raise HTTPException(503, "PDF conversion unavailable — Gotenberg not configured")
            try:
                import httpx
                from pathlib import Path as _Path
                from src.models.storage import StorageConfig
                from src.services.storage import get_storage_backend as _get_backend

                storage_cfg = await db.get(StorageConfig, rendered_doc.storage_config_id)
                backend = _get_backend(storage_cfg)
                docx_bytes = await backend.read(rendered_doc.file_path)

                async with httpx.AsyncClient(timeout=120) as client:
                    resp = await client.post(
                        f"{settings.gotenberg_url}/forms/libreoffice/convert",
                        files={"files": (out_filename, docx_bytes, rendered_doc.mime_type)},
                    )
                if resp.status_code != 200:
                    raise HTTPException(502, "PDF conversion failed")

                pdf_bytes = resp.content
                src_stem = _Path(out_filename).stem
                pdf_filename = body.output_filename.replace(".docx", ".pdf") if body.output_filename else f"{src_stem}.pdf"
                pdf_result = await backend.save(pdf_bytes, pdf_filename)

                pdf_doc = Document(
                    title=rendered_doc.title.replace(".docx", ".pdf") if ".docx" in rendered_doc.title else rendered_doc.title,
                    description=rendered_doc.description,
                    file_name=pdf_result.file_name,
                    original_filename=pdf_filename,
                    file_path=pdf_result.file_path,
                    file_size=pdf_result.file_size,
                    mime_type="application/pdf",
                    extension="pdf",
                    checksum=pdf_result.checksum,
                    storage_config_id=rendered_doc.storage_config_id,
                    owner_id=current_user.id,
                    folder_id=rendered_doc.folder_id,
                    source_type="generated",
                    source_metadata={
                        **(rendered_doc.source_metadata or {}),
                        "converted_from_docx_id": str(rendered_doc.id),
                    },
                )
                db.add(pdf_doc)
                await db.flush()
                await db.refresh(pdf_doc)

                # Remove the intermediate DOCX (keep only PDF)
                rendered_doc.deleted_at = __import__("datetime").datetime.now(__import__("datetime").timezone.utc)
                final_document_id = pdf_doc.id

            except HTTPException:
                raise
            except Exception as e:
                raise HTTPException(502, f"PDF conversion error: {e}") from e

        session = await svc.update_session(session_id, {
            "status": "completed",
            "document_id": final_document_id,
            "folder_id": target_folder_id,
        })
        # Phase 3 — COMMIT CỐ Ý (status-machine): nhánh except ghi status="failed" RỒI
        # commit TRƯỚC `raise` → phải bền qua rollback của boundary. KHÔNG gỡ.
        await db.commit()
    except HTTPException:
        await svc.update_session(session_id, {"status": "failed", "error_message": "Generation failed"})
        await db.commit()
        raise

    return GeneratorSessionResponse.model_validate(session)


@router.delete("/sessions/{session_id}", status_code=204)
async def delete_session(
    session_id: uuid.UUID,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    svc = GeneratorService(db)
    session = await svc.get_session_for_user(session_id, current_user.id)
    if session is None:
        raise HTTPException(404, "Session not found")
    await svc.delete_session(session_id)
    # Phase 3: commit ở boundary (get_db).
    await db.flush()


# ─── Generator Session Versions (chỉnh sửa tay) ─────────────────────────────────

@router.post("/sessions/{session_id}/versions", response_model=GeneratorSessionVersionResponse)
async def create_session_version(
    session_id: uuid.UUID,
    body: GeneratorSessionVersionCreateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Lưu một bản chỉnh sửa tay → tạo version + cập nhật con trỏ edited_html của session."""
    from datetime import datetime, timezone

    svc = GeneratorService(db)
    session = await svc.get_session_for_user(session_id, current_user.id)
    if session is None:
        raise HTTPException(404, "Session not found")

    version_no = await svc.next_version_no(session_id)
    label = body.label or f"V{version_no} — {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')}"

    version = await svc.create_version({
        "session_id": session_id,
        "version_no": version_no,
        "label": label,
        "edited_html": body.edited_html,
        "field_values": body.field_values or {},
    })
    # Cập nhật con trỏ "đang làm việc" của session về version mới nhất
    update: dict = {"edited_html": body.edited_html}
    if body.field_values:
        update["field_values"] = body.field_values
    await svc.update_session(session_id, update)
    # Phase 3: commit ở boundary (get_db).
    await db.flush()
    await db.refresh(version)
    return GeneratorSessionVersionResponse.model_validate(version)


@router.get("/sessions/{session_id}/versions", response_model=GeneratorSessionVersionListResponse)
async def list_session_versions(
    session_id: uuid.UUID,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
    limit: int = Query(50, ge=1, le=100),
    offset: int = Query(0, ge=0),
):
    svc = GeneratorService(db)
    session = await svc.get_session_for_user(session_id, current_user.id)
    if session is None:
        raise HTTPException(404, "Session not found")

    items, total = await svc.list_versions_for_session(session_id, limit=limit, offset=offset)
    return GeneratorSessionVersionListResponse(
        items=[GeneratorSessionVersionResponse.model_validate(v) for v in items],
        total=total,
    )


@router.patch("/sessions/{session_id}/versions/{version_id}", response_model=GeneratorSessionVersionResponse)
async def update_session_version(
    session_id: uuid.UUID,
    version_id: uuid.UUID,
    body: GeneratorSessionVersionUpdateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Đổi tên (label) một version."""
    svc = GeneratorService(db)
    session = await svc.get_session_for_user(session_id, current_user.id)
    if session is None:
        raise HTTPException(404, "Session not found")

    version = await svc.get_version_for_session(version_id, session_id)
    if version is None:
        raise HTTPException(404, "Version not found")

    if body.label is not None:
        version = await svc.update_version(version_id, {"label": body.label})
    # Phase 3: commit ở boundary (get_db).
    await db.flush()
    await db.refresh(version)
    return GeneratorSessionVersionResponse.model_validate(version)


@router.delete("/sessions/{session_id}/versions/{version_id}", status_code=204)
async def delete_session_version(
    session_id: uuid.UUID,
    version_id: uuid.UUID,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Xoá một version."""
    svc = GeneratorService(db)
    session = await svc.get_session_for_user(session_id, current_user.id)
    if session is None:
        raise HTTPException(404, "Session not found")

    version = await svc.get_version_for_session(version_id, session_id)
    if version is None:
        raise HTTPException(404, "Version not found")
    await svc.delete_version(version_id)
    # Phase 3: commit ở boundary (get_db).
    await db.flush()


# ─── AI đề xuất chỉnh sửa (block-based) ─────────────────────────────────────────
# Business-logic AI-revise (_AI_REVISE_SYSTEM/_strip_json_fence/propose_ops) đã trích
# sang GeneratorService (Phase 8 A3b). Helper stream-disconnect giữ ở route (HTTP-bound).


async def _consume_stream_with_disconnect(stream, request) -> str:
    """Tiêu thụ async stream text-delta, RAISE HTTPException(499) nếu client ngắt
    giữa chừng (Phase 4 T5b — tách module-level để test được).

    Disconnect được kiểm tra mỗi lần nhận một text-delta (``gateway.stream`` chỉ
    yield delta khác rỗng). Khi generate, content token về liên tục nên 499 vẫn
    bắt kịp thời (chỉ không check trên các chunk rỗng cuối stream — lúc đó đã
    generate xong). Raise trong ``async for`` đẩy GeneratorExit lan ngược chuỗi
    generator → provider connection được dọn theo cơ chế cleanup của asyncio
    (giống hành vi code cũ trước khi qua gateway; KHÔNG phải cam kết hủy tức thì)."""
    accumulated = ""
    async for delta in stream:
        if await request.is_disconnected():
            raise HTTPException(499, "Client disconnected")
        accumulated += delta
    return accumulated


@router.post("/sessions/{session_id}/ai-revise", response_model=AiReviseResponse)
async def ai_revise(
    session_id: uuid.UUID,
    body: AiReviseRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """AI đề xuất chỉnh sửa NỘI DUNG trên HTML đang làm việc (mức a).

    LLM chỉ thao tác ở mức TEXT theo block-id (Tuyến 1); backend áp dụng giữ
    nguyên thẻ + sanitize/guard (Tuyến 2). Không tự lưu version — FE quyết định
    duyệt rồi gọi /apply + tạo version.
    """
    svc = GeneratorService(db)
    session = await svc.get_session_for_user(session_id, current_user.id)
    if session is None:
        raise HTTPException(404, "Session not found")

    _, _, text_map = _html_blocks(body.html)
    if not text_map:
        raise HTTPException(422, "Tài liệu chưa có nội dung để chỉnh sửa.")

    valid_ids = set(text_map.keys())
    ops, warnings = await svc.propose_ops(text_map, body.instructions)

    # Lọc op trỏ tới block không tồn tại (Tuyến 2).
    kept: list[BlockEditOp] = []
    for op in ops:
        if op.block_id not in valid_ids:
            warnings.append(f"Bỏ qua thao tác tới block không tồn tại ({op.block_id}).")
            continue
        # Điền text gốc để FE hiển thị diff (replace/delete có before; insert thì không).
        if op.op in ("replace", "delete"):
            op.before_text = text_map.get(op.block_id)
        kept.append(op)

    applied_html, apply_warnings = _apply_block_ops(body.html, kept)
    revised_html_all = _sanitize_html(applied_html)
    diff_html = _build_diff_html(body.html, kept)
    warnings.extend(apply_warnings)
    warnings.extend(_structural_guard(body.html, revised_html_all))
    lost = _tokens_in(body.html) - _tokens_in(revised_html_all)
    if lost:
        warnings.append("Cảnh báo: một số trường bị mất sau chỉnh sửa: " + ", ".join(sorted(lost)))

    return AiReviseResponse(
        ops=kept,
        revised_html_all=revised_html_all,
        diff_html=diff_html,
        warnings=warnings,
    )


@router.post("/sessions/{session_id}/ai-revise/apply", response_model=AiReviseApplyResponse)
async def ai_revise_apply(
    session_id: uuid.UUID,
    body: AiReviseApplyRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Áp dụng các op đã được user DUYỆT lên HTML (tất định) → trả HTML cuối.

    Dùng chung lõi với /ai-revise (apply + sanitize + guard). FE gọi sau khi
    accept một phần / toàn bộ, rồi tự tạo version từ kết quả.
    """
    svc = GeneratorService(db)
    session = await svc.get_session_for_user(session_id, current_user.id)
    if session is None:
        raise HTTPException(404, "Session not found")

    applied_html, warnings = _apply_block_ops(body.html, body.ops)
    revised_html = _sanitize_html(applied_html)
    warnings.extend(_structural_guard(body.html, revised_html))
    lost = _tokens_in(body.html) - _tokens_in(revised_html)
    if lost:
        warnings.append("Cảnh báo: một số trường bị mất sau chỉnh sửa: " + ", ".join(sorted(lost)))

    # Round-trip check (không chặn) — đảm bảo HTML cuối còn xuất được DOCX.
    try:
        _html_to_docx_bytes(revised_html)
    except Exception:
        warnings.append("Cảnh báo: nội dung có thể không xuất DOCX hoàn hảo.")

    return AiReviseApplyResponse(revised_html=revised_html, warnings=warnings)


# ─── Document to Template ──────────────────────────────────────────────────────

@router.post("/document-to-template", response_model=DocumentToTemplateResponse)
async def document_to_template(
    request: Request,
    body: DocumentToTemplateRequest,
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
):
    """Convert một document thành template, TÁI DÙNG pipeline trích của Beta
    (`template_service`): location-based + replace ở mức run → GIỮ NGUYÊN FORMAT
    gốc (font, bảng, heading, header/footer). Có chunking cho tài liệu dài.
    """
    from src.ai import AIGateway
    from src.schemas.template import TemplateFieldResponse
    # Dùng BẢN SAO RIÊNG của generator (duplicate luồng Beta) — không đụng template_service.py
    from src.services.generator_template_service import commit_template, extract_template_draft

    source_doc = await db.get(Document, body.document_id)
    if source_doc is None or source_doc.deleted_at is not None:
        raise HTTPException(404, "Document not found")
    # Chống IDOR: chỉ user có quyền viewer mới chuyển document thành template.
    from src.services.document_permission import DocumentPermissionService
    await DocumentPermissionService(db).check_permission(
        current_user, document_id=source_doc.id, required="viewer"
    )

    if source_doc.source_type == "template":
        # Đã là template — trả thông tin sẵn có, không gọi AI
        meta = source_doc.source_metadata or {}
        fields = meta.get("template_fields") or []
        return DocumentToTemplateResponse(
            template_id=source_doc.id,
            title=source_doc.title,
            field_count=len(fields),
            detected_count=len(fields),
            template_fields=[TemplateFieldResponse(**f) for f in fields],
        )

    # Engine Beta dùng python-docx → chỉ DOCX. PDF sẽ hỗ trợ sau.
    ext = Path(source_doc.original_filename).suffix.lower()
    if ext not in (".docx", ".doc"):
        raise HTTPException(
            422,
            "Hiện chỉ hỗ trợ tạo template giữ định dạng từ file DOCX. PDF sẽ được hỗ trợ sau.",
        )

    # Checkpoint ①: trước khi gọi LLM
    if await request.is_disconnected():
        raise HTTPException(499, "Client disconnected")

    # llm_call streaming — cùng signature với non-streaming, service không cần biết.
    # Streaming cho phép Azure dừng generate ngay khi client ngắt → tiết kiệm usage.
    gw = AIGateway(db)

    async def llm_call(messages, response_format=None):
        overrides = {"response_format": response_format} if response_format else {}
        return await _consume_stream_with_disconnect(
            gw.stream(messages, **overrides), request
        )

    # 1) Trích field theo location — không side effect
    draft = await extract_template_draft(db, body.document_id, llm_call=llm_call)
    if draft.get("error"):
        raise HTTPException(422, draft["error"])
    draft_fields = draft.get("fields", []) or []
    detected_count = len(draft_fields)

    # Checkpoint ②: sau LLM, trước DB write — phòng trường hợp ngắt cuối stream
    if await request.is_disconnected():
        raise HTTPException(499, "Client disconnected")

    # 2) Commit: thay tại chỗ trên docx GỐC (giữ format) + tạo Document template
    committed = await commit_template(
        db,
        body.document_id,
        draft_fields,
        current_user.id,
        draft.get("doc_description"),
    )
    if committed.get("error"):
        raise HTTPException(422, committed["error"])

    template_id = uuid.UUID(committed["template_id"])

    # 3) Override title nếu người dùng cung cấp (Beta tự đặt "X (template)")
    tmpl = await db.get(Document, template_id)
    if body.title and tmpl:
        tmpl.title = body.title
        # Phase 3: commit ở boundary (get_db).
        await db.flush()
        await db.refresh(tmpl)
    title = tmpl.title if tmpl else (body.title or source_doc.title)

    template_fields = committed.get("template_fields", [])
    return DocumentToTemplateResponse(
        template_id=template_id,
        title=title,
        field_count=committed.get("replaced_count", len(template_fields)),
        detected_count=detected_count,
        template_fields=[TemplateFieldResponse(**f) for f in template_fields],
    )


@router.get("/history")
async def get_history(
    current_user: CurrentUser,
    db: AsyncSession = Depends(get_db),
    limit: int = Query(20, ge=1, le=100),
    offset: int = Query(0, ge=0),
):
    query = (
        select(Document)
        .where(
            Document.source_type == "generated",
            Document.owner_id == current_user.id,
            Document.deleted_at.is_(None),
        )
        .order_by(Document.created_at.desc())
        .limit(limit)
        .offset(offset)
    )
    result = await db.execute(query)
    docs = list(result.scalars().all())
    return {
        "items": [
            {
                "id": str(d.id),
                "title": d.title,
                "original_filename": d.original_filename,
                "created_at": d.created_at.isoformat(),
                "template_id": (d.source_metadata or {}).get("template_id"),
            }
            for d in docs
        ]
    }
