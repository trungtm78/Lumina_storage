"""Convert Got It Legal DOCX templates to the `{placeholder}` syntax.

The legal team's master DOCX files use a mix of fill-in markers:
  - `[ALLCAPS LABEL]` — bracket-wrapped Vietnamese caps (e.g. `[CÔNG TY]`)
  - `[*]` — generic "fill in" marker (number/code goes here)
  - `___/___/____` or `_____/______/_____` — underscore date triples

The Document Generator engine only recognises `{snake_case_name}` placeholders,
so we run a regex-based conversion in place that:
  1. Slugifies bracket-caps labels deterministically (so the same label always
     maps to the same placeholder name across files).
  2. Numbers `[*]` markers sequentially per document.
  3. Replaces the first underscore-date with `{ngay_ky}`, subsequent ones with
     `{ngay_2}`, `{ngay_3}`, …

Output: rewritten DOCX bytes + the list of unique placeholders introduced.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass

from docx import Document
from docx.text.paragraph import Paragraph

# `[CAPS LABEL]` — at least 2 chars, allow Vietnamese diacritics, spaces, slashes.
_BRACKET_CAPS_RE = re.compile(r"\[([\w\s/&\-]{2,40})\]", re.UNICODE)
# Underscore-date triple: a/b/c where each part is 3+ underscores.
_UNDERSCORE_DATE_RE = re.compile(r"_{3,}\s*/\s*_{3,}\s*/\s*_{3,}")
# Single `[*]` fill marker.
_BRACKET_STAR_RE = re.compile(r"\[\s*\*\s*\]")


@dataclass
class ConversionResult:
    docx_bytes: bytes
    placeholders: list[str]
    """Unique placeholder names introduced, in first-occurrence order."""


def _strip_diacritics(s: str) -> str:
    nfkd = unicodedata.normalize("NFD", s)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def slugify(label: str) -> str:
    """Turn a bracket-caps label into a snake_case placeholder name.

    Examples:
        "CÔNG TY"        → "cong_ty"
        "NĂM"            → "nam"
        "MÃ SỐ THUẾ"     → "ma_so_thue"
        "Address/Email"  → "address_email"
    """
    cleaned = _strip_diacritics(label).lower()
    # Replace any non-alphanumeric run with a single underscore.
    slug = re.sub(r"[^a-z0-9]+", "_", cleaned).strip("_")
    if not slug:
        slug = "field"
    # Placeholder must start with a letter (the rescan regex enforces this).
    if not slug[0].isalpha():
        slug = "f_" + slug
    return slug


def _process_text(
    text: str,
    bracket_caps_map: dict[str, str],
    counters: dict[str, int],
) -> tuple[str, list[str]]:
    """Apply all conversion rules to a single text fragment.

    `bracket_caps_map` lets us reuse the same placeholder for the same label
    across the document. `counters` tracks `star`/`date` sequence numbers
    across the whole document.
    """
    introduced: list[str] = []

    # 1. Bracket-caps labels — only match if the inside is mostly uppercase
    #    (avoids replacing inline references like `[note 1]`).
    def _bracket_caps_sub(m: re.Match) -> str:
        inner = m.group(1).strip()
        # Heuristic: at least one alpha char, and uppercase letters dominate.
        alpha = [c for c in inner if c.isalpha()]
        if not alpha:
            return m.group(0)
        upper_ratio = sum(1 for c in alpha if c.isupper()) / len(alpha)
        if upper_ratio < 0.6:
            return m.group(0)
        # Re-use the same slug for the same label — keep template fields stable.
        if inner not in bracket_caps_map:
            slug = slugify(inner)
            # Disambiguate if slug collides with an already-used different label.
            existing = set(bracket_caps_map.values())
            base = slug
            n = 2
            while slug in existing:
                slug = f"{base}_{n}"
                n += 1
            bracket_caps_map[inner] = slug
            introduced.append(slug)
        return "{" + bracket_caps_map[inner] + "}"

    text = _BRACKET_CAPS_RE.sub(_bracket_caps_sub, text)

    # 2. Underscore-date triples → `{ngay_ky}`, `{ngay_2}`, …
    def _date_sub(_m: re.Match) -> str:
        counters["date"] = counters.get("date", 0) + 1
        n = counters["date"]
        slug = "ngay_ky" if n == 1 else f"ngay_{n}"
        introduced.append(slug)
        return "{" + slug + "}"

    text = _UNDERSCORE_DATE_RE.sub(_date_sub, text)

    # 3. `[*]` markers → `{so_hd}`, `{field_2}`, …
    def _star_sub(_m: re.Match) -> str:
        counters["star"] = counters.get("star", 0) + 1
        n = counters["star"]
        slug = "so_hd" if n == 1 else f"field_{n}"
        introduced.append(slug)
        return "{" + slug + "}"

    text = _BRACKET_STAR_RE.sub(_star_sub, text)

    return text, introduced


def _rewrite_paragraph_text(para: Paragraph, new_text: str) -> None:
    """Replace a paragraph's text while keeping run 0's formatting.

    DOCX runs may split a single placeholder across multiple chunks of text
    (e.g. `[CÔNG` and ` TY]` in separate runs). The safe rewrite is to merge
    all text into the first run and clear the rest.
    """
    if not para.runs:
        return
    para.runs[0].text = new_text
    for run in para.runs[1:]:
        run.text = ""


def _iter_all_paragraphs(doc: Document):
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para
    for section in doc.sections:
        for part in (section.header, section.footer):
            if part is not None:
                for para in part.paragraphs:
                    yield para


def convert_docx(docx_bytes: bytes) -> ConversionResult:
    """Run all placeholder rewrites on a DOCX and return the rewritten bytes."""
    doc = Document(io.BytesIO(docx_bytes))

    bracket_caps_map: dict[str, str] = {}
    counters: dict[str, int] = {}
    seen_placeholders: list[str] = []
    seen_set: set[str] = set()

    for para in _iter_all_paragraphs(doc):
        original = para.text
        if not original:
            continue
        new_text, introduced = _process_text(original, bracket_caps_map, counters)
        if new_text != original:
            _rewrite_paragraph_text(para, new_text)
        for slug in introduced:
            if slug not in seen_set:
                seen_set.add(slug)
                seen_placeholders.append(slug)

    out = io.BytesIO()
    doc.save(out)
    return ConversionResult(
        docx_bytes=out.getvalue(),
        placeholders=seen_placeholders,
    )
