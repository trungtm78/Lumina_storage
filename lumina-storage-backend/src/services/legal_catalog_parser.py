"""Parse the GotIt Legal catalog xlsx into structured records.

The catalog is the single source of truth for which Sales B2B templates have
been approved by Legal. Each row is a logical template that may have a
Vietnamese-only and a bilingual VN-EN variant. We expand each row into one
record per existing language variant — entries marked "Chưa ban hành" (not
yet released) are skipped for that variant.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Iterable, Literal

from openpyxl import load_workbook

# Sheet 2 holds the catalog; sheet 1 is just structural notes.
CATALOG_SHEET_NAME = "2. Hướng dẫn áp dụng"

# Header rows occupy rows 1-3 (row 1 is title, row 2 is headers, row 3 is the
# language sub-header). Data rows start at row 4.
DATA_START_ROW = 4

LanguageVariant = Literal["single", "bilingual"]


@dataclass(frozen=True)
class CatalogEntry:
    """One template variant from the catalog (VIE OR BIL — never both)."""

    stt: int
    doc_category: str
    """LOẠI VĂN BẢN — e.g. 'Hợp đồng nguyên tắc', 'Phụ lục sản phẩm/dịch vụ'."""
    name: str
    """TÊN VĂN BẢN — the human-readable template name."""
    guidance: str
    """HƯỚNG DẪN ÁP DỤNG — when to use this template."""
    language: LanguageVariant
    """`single` for VIE-only, `bilingual` for VN-EN side-by-side."""
    file_code: str
    """The filename reference in the catalog (e.g. '01 - SAL_HOP DONG NGUYEN TAC_VIE')."""


_NOT_RELEASED_MARKERS = {"chưa ban hành", "chua ban hanh"}


def _is_released(cell: object) -> bool:
    if cell is None:
        return False
    s = str(cell).strip()
    if not s:
        return False
    return s.lower() not in _NOT_RELEASED_MARKERS


def parse_legal_catalog(xlsx_bytes: bytes) -> list[CatalogEntry]:
    """Read the xlsx bytes and return one CatalogEntry per existing variant.

    Empty rows, missing required columns, and not-yet-released variants are
    silently skipped. Rows where both variants are missing produce no entries.
    """
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    if CATALOG_SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"Catalog xlsx thiếu sheet '{CATALOG_SHEET_NAME}'. "
            f"Available: {wb.sheetnames}"
        )
    ws = wb[CATALOG_SHEET_NAME]

    entries: list[CatalogEntry] = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx < DATA_START_ROW:
            continue

        # Columns A..F: STT | LOẠI VĂN BẢN | TÊN VĂN BẢN | HƯỚNG DẪN | LINK_VIE | LINK_BIL
        if len(row) < 6:
            continue
        stt_cell, category, name, guidance, link_vie, link_bil = row[:6]

        # Skip empty / footer rows (no STT and no category)
        try:
            stt = int(stt_cell) if stt_cell is not None else None
        except (TypeError, ValueError):
            stt = None
        if stt is None or not category or not name:
            continue

        category_s = str(category).strip()
        name_s = str(name).strip()
        guidance_s = (str(guidance).strip() if guidance is not None else "")

        if _is_released(link_vie):
            entries.append(
                CatalogEntry(
                    stt=stt,
                    doc_category=category_s,
                    name=name_s,
                    guidance=guidance_s,
                    language="single",
                    file_code=str(link_vie).strip(),
                )
            )
        if _is_released(link_bil):
            entries.append(
                CatalogEntry(
                    stt=stt,
                    doc_category=category_s,
                    name=name_s,
                    guidance=guidance_s,
                    language="bilingual",
                    file_code=str(link_bil).strip(),
                )
            )

    wb.close()
    return entries


def categories_in_order(entries: Iterable[CatalogEntry]) -> list[str]:
    """Distinct LOẠI VĂN BẢN values in first-occurrence order."""
    seen: list[str] = []
    seen_set: set[str] = set()
    for e in entries:
        if e.doc_category not in seen_set:
            seen.append(e.doc_category)
            seen_set.add(e.doc_category)
    return seen
