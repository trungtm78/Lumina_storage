"""Báo cáo Word hoàn thành Phase 7 (God-service cleanup: agent.py + chat_service.py).

Chạy: MSYS_NO_PATHCONV=1 docker compose run --rm -v "C:/Lumina_Storage/docs:/out" \
  test python /app/tests/_gen_milestone7_report.py
"""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def para(t, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.italic = italic
    return p


def bullet(t, prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if prefix:
        p.add_run(prefix).bold = True
    p.add_run(t)
    return p


doc.add_heading("Lumina Storage — Báo cáo Hoàn thành Phase 7", level=0)
s = para("God-service cleanup: agent.py (sa_text→repo, write_file→service+ACL, gỡ except:pass) "
         "+ bắt đầu tách chat_service.py", italic=True)
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
m = para("Branch: refactor/milestone-1-foundation | Ngày: 2026-06-30 | 375 test PASS | import-linter 2 kept 0 broken")
m.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("1. Tổng quan & Mục tiêu", 1)
para(
    "agent.py (752 dòng) là god-service: truy vấn SQL THẲNG (sa_text) bỏ qua repository, write_file "
    "tự dựng Document (không ACL, không qua DocumentService), nhiều except: pass NUỐT lỗi im lặng. "
    "chat_service.py (782 dòng) trộn 5 trách nhiệm + còn hàm CHẾT stream_answer (~142 dòng). Phase 7: "
    "(a) sa_text → repo/service; (b) write_file → DocumentService.create_from_bytes + ACL; (c) gỡ "
    "except:pass (log BROAD); (d) bắt đầu tách chat_service — XÓA dead stream_answer + trích "
    "CitationService. Pure refactor: hành vi KHÔNG đổi (trừ giờ LOG); full suite + import-linter xanh."
)
bullet("agent.py: 0 sa_text, 0 except:pass; write_file qua service+ACL; chat_service 782→599 dòng.", prefix="Kết quả: ")
bullet("375 test PASS (+6 test mới); import-linter 2 kept 0 broken; outside-voice critique resolve 2 P0+P1.", prefix="Chất lượng: ")

doc.add_heading("2. Chi tiết thay đổi & Lý do (theo task)", 1)
tasks = [
    ("T1 — sa_text → repo (documents owner-only + systemconfig) (515fe7d)",
     "agent.py 4 sa_text bỏ qua repo → repo: DocumentRepository +3 method (search_workspace_files "
     "exclude skill_temp; list_workspace_files exclude skill_temp+template + extension dual; "
     "search_owned_templates source_type=template + ILIKE description) + SystemConfigRepository.get_by_key.",
     "Critique P0#1: 3 query KHÁC semantics → 3 method RIÊNG (không gộp). GIỮ owner-only chính xác "
     "(characterization test/method) — KHÔNG đổi sang get_accessible_paginated (sẽ broaden ACL)."),
    ("T2 — write_file → DocumentService.create_from_bytes + ACL (b105fbd)",
     "write_file tự dựng Document (raw storage + KHÔNG ACL) → DocumentService.create_from_bytes "
     "(no-commit, ACL check_permission khi có folder_id, storage qua _resolve_storage). write_file "
     "gọi service + giữ commit CỐ Ý 1 lần. Bỏ sa_text storage + import.",
     "Critique P0#2 (documented): LIMIT-1→get_default = nhất quán với upload thật (mục đích Phase 7; "
     "production luôn có default). P1#3: no double-commit (create_from_bytes no-commit)."),
    ("T3 — gỡ except:pass → log BROAD (8132511)",
     "4 block nuốt lỗi im lặng (run_script skill_model_config/vendor config/app_config + write_file) "
     "→ thêm logging. + import logging + logger module-level.",
     "Critique P1#7: GIỮ except Exception BROAD (thu hẹp type → exception khác escape = crash mới); "
     "chỉ thay pass bằng logger.warning(exc_info) + fall-through. Control-flow KHÔNG đổi."),
    ("T4 — xóa dead stream_answer + trích CitationService (cae0d6c)",
     "XÓA dead stream_answer (~142 dòng, simple LCEL pre-agent; route dùng stream_agent — 0 caller). "
     "Trích CitationService (_coerce_uuid/citation_to_source/citations_to_sources, pure transformation, "
     "0 phụ thuộc ChatService). chat_service 782→599 dòng.",
     "Critique: extraction an toàn NHẤT (test_citation_phase4 phủ 6 test, 0 circular). stream_answer "
     "xác minh dead (grep src/+tests/) trước khi xóa."),
    ("T5 — Checkpoint (commit này)",
     "Full suite 375 PASS; smoke (app boot 194 routes + agent/citation/document import); lint-imports "
     "2 kept 0 broken; grep agent.py 0 sa_text/0 except:pass; báo cáo Word.",
     "Khoá trạng thái sạch."),
]
for name, change, why in tasks:
    doc.add_heading(name, 2)
    p = doc.add_paragraph(); p.add_run("Thay đổi: ").bold = True; p.add_run(change)
    p = doc.add_paragraph(); p.add_run("Lý do: ").bold = True; p.add_run(why)

doc.add_heading("3. Đối chiếu Spec (alignment Phase 7)", 1)
para("Done-criteria spec §Phase 7 đạt đủ:")
bullet("Dọn agent.py: thay sa_text (storage_storageconfig/core_systemconfig + documents) bằng repo/service.")
bullet("Tách write_file god-function → qua DocumentService + ACL.")
bullet("Gỡ except:pass (giờ log, không nuốt im lặng).")
bullet("Tách dần chat_service.py: XÓA dead stream_answer + trích CitationService (bước đầu).")
para("Deviations có cơ sở (chốt từ outside-voice critique):", bold=True)
bullet("write_file storage LIMIT-1→get_default: behavior alignment (nhất quán upload thật; production luôn có default) — documented, không phải zero-change tuyệt đối.")
bullet("except:pass GIỮ catch BROAD + log (không thu hẹp type) → tránh crash mới; mục tiêu spec = hết SILENT swallow, đạt.")
bullet("chat_service mới tách CitationService (Concern C an toàn nhất); Concern B/D/E (ChatPermission/Title/stream_agent) + tách review.py/generator.py route = Phase 8.")

doc.add_heading("4. Kiểm thử & Smoke", 1)
bullet("Full suite 375 PASS (+6 test: test_agent_repo_cleanup×4, test_document_create_from_bytes×2; citation 6 chuyển sang citation_service).")
bullet("import-linter: 2 contract KEPT, 0 broken (agent.py/citation_service không phá layering).")
bullet("Smoke: app boot 194 routes; agent (build_model/create_agent) + citation_service + DocumentService import OK.")
bullet("grep agent.py: ZERO sa_text, ZERO except:pass; chat_service.py: ZERO stream_answer.")

doc.add_heading("5. Rủi ro & Backlog (Phase 8)", 1)
bullet("Tách nốt chat_service Concern B (ChatPermissionService) / D (TitleService) / E (stream_agent orchestration).")
bullet("Tách review.py (3734) / generator.py route (2328) → service+ai (router chỉ HTTP), từng endpoint.")
bullet("Module hóa domain/<x>/ (git mv từng domain 1 PR); worker DLQ + sửa race dedupe nuốt re-ingest + correlation job.")
bullet("7 @tool (write_file/list_directory/...) decorated nhưng chưa wire vào TOOLS — design question, giữ nguyên.")

doc.add_heading("6. Kết luận", 1)
para(
    "Phase 7 hoàn thành T1-T5 theo quy trình STRICT + outside-voice critique (resolve 2 P0 + P1 TRƯỚC "
    "khi code). agent.py sạch: 0 sa_text (qua repo/service), write_file qua DocumentService + ACL, 0 "
    "except:pass (giờ log BROAD — không crash mới). chat_service bắt đầu tách: xóa ~142 dòng dead + "
    "CitationService (782→599 dòng). 375 test PASS, import-linter 2 kept 0 broken, pure refactor "
    "(behavior alignment write_file storage documented).",
    bold=True,
)

out = "/out/Lumina_Storage_Milestone7_Bao_cao_Hoan_thanh.docx"
doc.save(out)
print("Saved:", out)
