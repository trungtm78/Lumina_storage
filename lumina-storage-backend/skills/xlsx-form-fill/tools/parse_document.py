"""Parse a .xlsx template and return its text content + detected blank cells.

Auto-called by the backend before each skill conversation turn — NOT an LLM tool.
"""

from __future__ import annotations

import re
from io import BytesIO


_BLANK_RE = re.compile(r'(\{{2}\w[\w\s]*\}{2})|([.…]{3,})|(_{3,})|(-{3,})', re.UNICODE)


def _is_blank(value: str) -> bool:
    return bool(_BLANK_RE.search(str(value)))


async def run(args: dict, ctx) -> dict:
    """Return sheet content as text and a summary of blank cells."""
    import openpyxl

    doc_bytes = await ctx.get_document_bytes(args["document_id"])
    wb = openpyxl.load_workbook(BytesIO(doc_bytes), data_only=True)

    lines: list[str] = []
    blank_fields: list[str] = []

    for sheet in wb.worksheets:
        lines.append(f"=== Sheet: {sheet.title} ===")
        for row in sheet.iter_rows():
            row_parts: list[str] = []
            for cell in row:
                val = cell.value
                if val is None:
                    continue
                val_str = str(val).strip()
                if val_str:
                    row_parts.append(val_str)
                    if _is_blank(val_str):
                        # Try to find label from left neighbour
                        label = val_str
                        if cell.column > 1:
                            left = sheet.cell(row=cell.row, column=cell.column - 1)
                            if left.value:
                                label = f"{left.value}: {val_str}"
                        blank_fields.append(label)
            if row_parts:
                lines.append("  ".join(row_parts))

    full_text = "\n".join(lines)

    if blank_fields:
        structure_lines = [f"Phát hiện {len(blank_fields)} ô cần điền:"]
        for f in blank_fields:
            structure_lines.append(f"  - {f}")
        structure = "\n".join(structure_lines)
    else:
        structure = "Không phát hiện ô trống rõ ràng. User có thể yêu cầu chỉnh sửa bất kỳ ô nào."

    return {
        "full_text": full_text[:8000],
        "structure": structure,
    }
