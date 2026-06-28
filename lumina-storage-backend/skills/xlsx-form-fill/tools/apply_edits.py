"""Apply a list of cell edits to a .xlsx template and save a rendered copy.

Supports two actions:
  replace    — replace an existing cell value (default)
  insert_row — insert a new row after a reference row

Edit format for replace:
  {context, old, new}

Edit format for insert_row:
  {action: "insert_row", context: "Sheet: ...", after: "cell value to insert after", values: [...]}

Always applies to the ORIGINAL template document.
"""

from __future__ import annotations

import re
from io import BytesIO
from copy import copy

_BLANK_PAT = re.compile(r'[.…_\-]{3,}', re.UNICODE)
_NORM_BLANK = re.compile(r'[.…_\-]{2,}', re.UNICODE)


def _normalise(text: str) -> str:
    return _NORM_BLANK.sub('___', str(text)).strip()


# ── Sheet / context helpers ────────────────────────────────────────────


def _sheet_matches_context(sheet_title: str, context: str) -> bool:
    if not context:
        return True
    ctx_lower = context.lower()
    title_lower = sheet_title.lower()
    return title_lower in ctx_lower or ctx_lower in title_lower


def _row_matches_context(row, context: str) -> bool:
    if not context:
        return True
    ctx_lower = context.lower()[:40]
    for cell in row:
        if cell.value is not None and ctx_lower in str(cell.value).lower():
            return True
    return False


def _left_label_matches(sheet, cell, context: str) -> bool:
    if not context or cell.column <= 1:
        return False
    left = sheet.cell(row=cell.row, column=cell.column - 1)
    if left.value is not None:
        left_str = str(left.value).strip().lower()
        ctx_lower = context.lower()[:40]
        return ctx_lower in left_str or left_str in ctx_lower
    return False


# ── Replace action ─────────────────────────────────────────────────────


def _apply_replace(wb, context: str, old: str, new: str) -> bool:
    """Replace a cell value. 4-pass strategy (exact → normalised → blank → fallback)."""
    if not old and not context:
        return False

    target_sheets = [s for s in wb.worksheets if _sheet_matches_context(s.title, context)]
    all_sheets = target_sheets if target_sheets else list(wb.worksheets)

    # Pass 1: Context + exact old
    for sheet in all_sheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if str(cell.value).strip() == old:
                    if not context or _row_matches_context(row, context) or _left_label_matches(sheet, cell, context):
                        cell.value = new
                        return True

    # Pass 2: Normalised blank match
    if old and _BLANK_PAT.search(old):
        norm_old = _normalise(old)
        for sheet in all_sheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    if _normalise(str(cell.value)) == norm_old:
                        if not context or _row_matches_context(row, context) or _left_label_matches(sheet, cell, context):
                            cell.value = new
                            return True

    # Pass 3: Context row → replace first blank cell
    if context:
        for sheet in all_sheets:
            for row in sheet.iter_rows():
                if not _row_matches_context(row, context):
                    continue
                for cell in row:
                    if cell.value is not None and _BLANK_PAT.search(str(cell.value)):
                        cell.value = new
                        return True
                for cell in row:
                    if cell.value is None and cell.column > 1:
                        left = sheet.cell(row=cell.row, column=cell.column - 1)
                        if left.value is not None and context.lower()[:20] in str(left.value).lower():
                            cell.value = new
                            return True

    # Pass 4: Exact old anywhere
    if old:
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None and str(cell.value).strip() == old:
                        cell.value = new
                        return True

    return False


# ── Insert row action ──────────────────────────────────────────────────


def _find_reference_row(sheet, after_value: str) -> int | None:
    """Find the row number that contains `after_value` in any cell."""
    after_lower = str(after_value).strip().lower()
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip().lower() == after_lower:
                return cell.row
    return None


def _find_last_data_row(sheet, start_col: int = 2) -> int:
    """Find the last row with data in the given column."""
    last = 1
    for row in sheet.iter_rows(min_col=start_col, max_col=start_col):
        for cell in row:
            if cell.value is not None and str(cell.value).strip() and str(cell.value).strip() != '–':
                last = max(last, cell.row)
    return last


def _copy_cell_style(src_cell, dst_cell):
    """Copy formatting from source cell to destination cell."""
    try:
        if src_cell.has_style:
            dst_cell.font = copy(src_cell.font)
            dst_cell.border = copy(src_cell.border)
            dst_cell.fill = copy(src_cell.fill)
            dst_cell.number_format = src_cell.number_format
            dst_cell.alignment = copy(src_cell.alignment)
    except Exception:
        pass


def _apply_insert_row(wb, context: str, after: str, values: list) -> bool:
    """Insert a new row with values after a reference row."""
    if not values:
        return False

    target_sheets = [s for s in wb.worksheets if _sheet_matches_context(s.title, context)]
    sheet = target_sheets[0] if target_sheets else wb.worksheets[0]

    # Find where to insert
    insert_after_row = None
    if after:
        insert_after_row = _find_reference_row(sheet, after)

    if insert_after_row is None:
        # Fallback: insert after last data row
        insert_after_row = _find_last_data_row(sheet)

    target_row = insert_after_row + 1

    # Shift rows down to make space
    sheet.insert_rows(target_row)

    # Determine starting column — use same as reference row
    # Find first non-empty column in reference row
    start_col = 1
    for cell in sheet[insert_after_row]:
        if cell.value is not None:
            start_col = cell.column
            break

    # Write values and copy style from row above
    for i, val in enumerate(values):
        dst = sheet.cell(row=target_row, column=start_col + i)
        # Try to convert numeric strings
        if isinstance(val, str):
            try:
                val = int(val)
            except ValueError:
                try:
                    val = float(val)
                except ValueError:
                    pass
        dst.value = val
        # Copy style from the cell above
        src = sheet.cell(row=insert_after_row, column=start_col + i)
        _copy_cell_style(src, dst)

    return True


# ── Main entry ─────────────────────────────────────────────────────────


async def run(args: dict, ctx) -> dict:
    """Apply edits to the xlsx template and save a rendered document."""
    import openpyxl

    document_id: str = args["document_id"]
    edits: list[dict] = args.get("edits", [])

    doc_bytes = await ctx.get_document_bytes(document_id)
    wb = openpyxl.load_workbook(BytesIO(doc_bytes))

    applied_count = 0
    failed_edits: list[dict] = []

    for edit in edits:
        action = edit.get("action", "replace")
        context = edit.get("context", "")

        if action == "insert_row":
            after = edit.get("after", "")
            values = edit.get("values", [])
            if _apply_insert_row(wb, context, after, values):
                applied_count += 1
            else:
                failed_edits.append(edit)
        else:
            old = str(edit.get("old", "")).strip()
            new = edit.get("new", "")
            if not old and not context:
                continue
            if _apply_replace(wb, context, old, new):
                applied_count += 1
            else:
                failed_edits.append(edit)

    output = BytesIO()
    wb.save(output)
    rendered_bytes = output.getvalue()

    rendered_id = await ctx.save_rendered_document(rendered_bytes, document_id)

    # Generate PDF preview via Gotenberg
    preview_pdf_id = None
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    pdf_bytes = await ctx.convert_to_pdf(rendered_bytes, mime)
    if pdf_bytes:
        preview_pdf_id = await ctx.save_pdf_preview(pdf_bytes, document_id)

    return {
        "rendered_document_id": str(rendered_id),
        "preview_pdf_id": str(preview_pdf_id) if preview_pdf_id else None,
        "applied_count": applied_count,
        "failed_edits": failed_edits,
    }
