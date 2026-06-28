import json

def cell(src):
    return {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [], "source": src}

def md(src):
    return {"cell_type": "markdown", "metadata": {}, "source": src}

cells = [
    md("# LLM-assisted Schema Parsing\n\nFlow: Excel upload → openpyxl extract raw → LLM parse schema → store mapping → Text-to-Pandas query"),

    # Cell 0: Config
    cell("""\
import asyncio, json, os
import asyncpg
import litellm

DB_DSN    = "postgresql://postgres:postgres123@192.168.1.229:5432/lumina_driver_dev"
XLSX_PATH = "uploads/2026/03/97fb555d-8c66-47b5-914b-b81f81c3d813.xlsx"

_PREFIXES = {"openai":"","azure":"azure/","anthropic":"anthropic/","google":"gemini/","ollama":"ollama/"}
def build_model(provider, model_name):
    p = _PREFIXES.get(provider, "")
    return model_name if (p and model_name.startswith(p)) else f"{p}{model_name}"

async def _load(purpose):
    conn = await asyncpg.connect(DB_DSN)
    row  = await conn.fetchrow(
        "SELECT provider,model_name,api_key,base_url,extra_config FROM core_aimodelconfig"
        " WHERE purpose=$1 AND is_default=TRUE AND is_active=TRUE LIMIT 1", purpose)
    await conn.close()
    raw   = row["extra_config"]
    extra = json.loads(raw) if raw and raw != "null" else {}
    return {"model": build_model(row["provider"], row["model_name"]),
            "api_key": row["api_key"], "api_base": row["base_url"], **extra}

chat_cfg = await _load("chat")
if "azure/" in chat_cfg["model"]:
    os.environ["AZURE_API_VERSION"] = "2024-12-01-preview"
print("Chat model:", chat_cfg["model"])
"""),

    # Cell 1: openpyxl helpers (border, bold, merged)
    cell("""\
import openpyxl
import pandas as pd
import numpy as np

def get_merged_value(ws, row, col, merged_ranges):
    for rng in merged_ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return None

def cell_has_border(cell):
    b = cell.border
    return any(getattr(b, s).style is not None for s in ("top","bottom","left","right"))

def find_header_start(ws, n_cols, max_check=30):
    for row_num in range(1, max_check + 1):
        count = sum(1 for c in range(1, min(n_cols+1,15))
                    if cell_has_border(ws.cell(row_num, c)))
        if count >= 3:
            return row_num
    return 1

def is_data_row_col_a(ws, row_num, merged_ranges):
    val = ws.cell(row_num, 1).value
    if val is None:
        val = get_merged_value(ws, row_num, 1, merged_ranges)
    if val is None: return False
    s = str(val).strip()
    if not s or s == "nan": return False
    try: float(s); return False
    except (ValueError, TypeError): pass
    if s == s.upper() and len(s) > 1: return False
    return True

def find_data_start(ws, merged_ranges):
    for row_num in range(1, ws.max_row + 1):
        if is_data_row_col_a(ws, row_num, merged_ranges):
            return row_num
    return 1

def fill_row_nearest(values):
    result = list(values)
    non_none = [i for i, v in enumerate(result) if v is not None]
    if not non_none: return result
    first, last = non_none[0], non_none[-1]
    for i in range(first, last + 1):
        if result[i] is None:
            left  = max((p for p in non_none if p < i), default=None)
            right = min((p for p in non_none if p > i), default=None)
            if left is not None and right is not None:
                result[i] = result[left] if (i-left) < (right-i) else result[right]
            elif left is not None: result[i] = result[left]
            elif right is not None: result[i] = result[right]
    return result

def get_actual_ncols(xlsx_path, sheet_name, nrows=20):
    df = pd.read_excel(xlsx_path, sheet_name=sheet_name, header=None, nrows=nrows)
    return df.dropna(axis=1, how="all").shape[1]

def build_headers(ws, header_start, data_start, n_cols, merged_ranges):
    row_filled = []
    for row_num in range(header_start, data_start):
        vals = []
        for col_num in range(1, n_cols + 1):
            v = ws.cell(row_num, col_num).value
            if v is None:
                v = get_merged_value(ws, row_num, col_num, merged_ranges)
            vals.append(str(v).strip() if v is not None else None)
        row_filled.append(fill_row_nearest(vals))
    headers = []
    for ci in range(n_cols):
        parts = []
        for rv in row_filled:
            v = rv[ci]
            if v and v not in ("nan","None") and v not in parts:
                parts.append(v)
        headers.append("-".join(parts) if parts else f"COL_{ci}")
    return headers

def load_sheet(wb, xlsx_path, sheet_name):
    ws = wb[sheet_name]
    merged = list(ws.merged_cells.ranges)
    n_cols = get_actual_ncols(xlsx_path, sheet_name)
    data_start  = find_data_start(ws, merged)
    header_start = find_header_start(ws, n_cols)
    headers = build_headers(ws, header_start, data_start, n_cols, merged)

    # Extract data rows with bold flag
    rows = []
    for row_num in range(data_start, ws.max_row + 1):
        row_data = []
        for col_num in range(1, n_cols + 1):
            v = ws.cell(row_num, col_num).value
            if v is None:
                v = get_merged_value(ws, row_num, col_num, merged)
            row_data.append(v)
        bold = ws.cell(row_num, 1).font.bold or False
        rows.append({"bold": bold, "data": row_data})

    df = pd.DataFrame([r["data"] for r in rows], columns=headers)
    df = df.dropna(axis=1, how="all")
    fc = df.columns[0]
    df = df[df[fc].notna() & df[fc].astype(str).str.strip().ne("")]
    df = df.reset_index(drop=True)

    bold_flags = [rows[i]["bold"] for i in range(len(rows))
                  if rows[i]["data"][0] is not None
                  and str(rows[i]["data"][0]).strip() not in ("", "nan")]
    bold_flags = bold_flags[:len(df)]

    return df, bold_flags, headers, header_start, data_start

wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
print("Sheets:", wb.sheetnames)
"""),

    # Cell 2: Build raw preview for LLM (1 sheet)
    cell("""\
def build_raw_preview(wb, xlsx_path, sheet_name, n_preview_rows=12):
    \"\"\"Trích xuất raw rows đầu tiên + tên cột raw để gửi LLM.\"\"\"\
    ws = wb[sheet_name]
    merged = list(ws.merged_cells.ranges)
    n_cols = get_actual_ncols(xlsx_path, sheet_name)
    data_start = find_data_start(ws, merged)
    header_start = find_header_start(ws, n_cols)
    headers = build_headers(ws, header_start, data_start, n_cols, merged)

    # First few data rows + bold info
    preview_rows = []
    bold_preview = []
    for row_num in range(data_start, min(data_start + n_preview_rows, ws.max_row + 1)):
        vals = []
        for col_num in range(1, n_cols + 1):
            v = ws.cell(row_num, col_num).value
            if v is None:
                v = get_merged_value(ws, row_num, col_num, merged)
            vals.append(str(v).strip() if v is not None else "")
        bold_preview.append(ws.cell(row_num, 1).font.bold or False)
        preview_rows.append(vals)

    # Format as table string (only first 12 cols để tránh quá dài)
    MAX_COLS_PREVIEW = 12
    col_headers = headers[:MAX_COLS_PREVIEW]
    lines = ["  ".join(f"{h[:18]:18s}" for h in col_headers)]
    lines.append("-" * (20 * min(MAX_COLS_PREVIEW, len(col_headers))))
    for i, (row, bold) in enumerate(zip(preview_rows, bold_preview)):
        prefix = "[BOLD] " if bold else "       "
        cells  = "  ".join(f"{str(v)[:18]:18s}" for v in row[:MAX_COLS_PREVIEW])
        lines.append(f"{prefix}{cells}")

    return {
        "sheet_name": sheet_name,
        "n_cols": n_cols,
        "headers": headers,
        "table_preview": "\\n".join(lines),
        "bold_flags_preview": bold_preview,
    }

preview = build_raw_preview(wb, XLSX_PATH, "DT 01.2026")
print(f"Sheet: {preview['sheet_name']}, {preview['n_cols']} cols")
print(f"\\nHeaders (first 12):")
for i, h in enumerate(preview["headers"][:12]):
    print(f"  [{i}] {h}")
print(f"\\nTable preview:")
print(preview["table_preview"])
"""),

    # Cell 3: LLM parse schema
    cell("""\
def llm_parse_schema(chat_cfg, preview):
    prompt = f\"\"\"Bạn là data analyst. Dưới đây là dữ liệu từ sheet Excel "{preview['sheet_name']}".
Dòng [BOLD] là dòng parent (ví dụ: tên cửa hàng), dòng không bold là sub-category.

Tên cột hiện tại (đã parse từ merged headers):
{json.dumps(preview['headers'][:20], ensure_ascii=False, indent=2)}

Preview 12 dòng đầu (12 cột đầu):
{preview['table_preview']}

Hãy trả về JSON với format sau (CHỈ JSON, không giải thích):
{{
  "sheet_description": "mô tả ngắn sheet này chứa gì",
  "row_structure": "mô tả cấu trúc row: parent là gì, sub-category là gì",
  "key_columns": {{
    "store_name": "tên cột chứa tên cửa hàng/entity chính",
    "category": "tên cột chứa sub-category (nếu có, else null)",
    "main_revenue": "tên cột chứa doanh thu kỳ này quan trọng nhất",
    "target": "tên cột chứa kế hoạch/target (nếu có, else null)",
    "prev_period": "tên cột doanh thu kỳ trước (nếu có, else null)"
  }},
  "parent_detection": "bold=True"
}}
\"\"\"
    kw   = {{k: v for k, v in chat_cfg.items() if k != "model" and v is not None}}
    resp = litellm.completion(
        model=chat_cfg["model"],
        messages=[{{"role":"user","content":prompt}}],
        temperature=0, **kw)
    raw = resp.choices[0].message.content.strip()
    # Strip markdown code block nếu có
    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    return json.loads(raw.strip())

schema = llm_parse_schema(chat_cfg, preview)
print(json.dumps(schema, ensure_ascii=False, indent=2))
"""),

    # Cell 4: Load full DataFrame + add store_name column
    cell("""\
def load_with_hierarchy(wb, xlsx_path, sheet_name, schema):
    \"\"\"Load DataFrame, thêm cột store_name từ bold parent rows.\"\"\"\
    df, bold_flags, headers, _, _ = load_sheet(wb, xlsx_path, sheet_name)

    # Propagate store name từ bold rows xuống sub-rows
    store_col = schema["key_columns"]["store_name"]
    cat_col   = schema["key_columns"].get("category")

    if store_col in df.columns:
        current_store = None
        store_names   = []
        categories    = []
        for i, (_, row) in enumerate(df.iterrows()):
            is_bold = bold_flags[i] if i < len(bold_flags) else False
            if is_bold:
                current_store = str(row[store_col]).strip()
                store_names.append(current_store)
                categories.append(None)          # parent row không có category
            else:
                store_names.append(current_store)
                cat_val = str(row[store_col]).strip() if cat_col is None else str(row.get(cat_col, row[store_col])).strip()
                categories.append(cat_val)

        df = df.copy()
        df.insert(0, "__store__",    store_names)
        df.insert(1, "__category__", categories)

    return df, bold_flags

df_full, bold_flags = load_with_hierarchy(wb, XLSX_PATH, "DT 01.2026", schema)
print(f"Shape: {df_full.shape}")
print(f"\\nFirst 8 rows (store + category + key cols):")
key_cols = ["__store__", "__category__"] + [schema["key_columns"]["main_revenue"],
             schema["key_columns"]["target"] or ""]
key_cols = [c for c in key_cols if c and c in df_full.columns]
df_full[key_cols].head(8)
"""),

    # Cell 5: Split parent vs sub-rows
    cell("""\
# Parent rows (store summary) — bold=True
df_stores = df_full[[b for b in [True]] and df_full.iloc[
    [i for i, b in enumerate(bold_flags) if b]
]].copy()

# Sub-category rows — bold=False
df_subs = df_full.iloc[
    [i for i, b in enumerate(bold_flags) if not b]
].copy()

print(f"Parent rows (stores):   {len(df_stores)}")
print(f"Sub-category rows:      {len(df_subs)}")
print(f"\\nCategories found: {sorted(df_subs['__category__'].dropna().unique().tolist())}")
"""),

    # Cell 6: Build chunks (RAGFlow style — key-value per row)
    cell("""\
def row_to_chunk(row, sheet_name, schema):
    rev_col  = schema["key_columns"]["main_revenue"]
    tgt_col  = schema["key_columns"]["target"]
    prev_col = schema["key_columns"]["prev_period"]

    lines = [f"[{sheet_name}]"]
    store = row.get("__store__", "")
    cat   = row.get("__category__", "")
    if store: lines.append(f"- Cửa hàng: {store}")
    if cat:   lines.append(f"- Nhóm: {cat}")

    # Key metrics
    for col in [rev_col, tgt_col, prev_col]:
        if col and col in row and pd.notna(row[col]):
            v = row[col]
            if isinstance(v, float) and v != 0:
                lines.append(f"- {col}: {v:,.0f}")
    return "\\n".join(lines)

# Demo: sub-rows của 98 Nguyễn Văn Trỗi
import unicodedata
def norm(s):
    return unicodedata.normalize("NFD", str(s)).encode("ascii","ignore").decode().lower()

mask = df_subs["__store__"].astype(str).map(norm).str.contains(norm("Nguyen Van Troi"))
sample = df_subs[mask]

print(f"Sub-rows of '98 Nguyễn Văn Trỗi': {len(sample)}")
for _, row in sample.iterrows():
    print()
    print(row_to_chunk(row.to_dict(), "DT 01.2026", schema))
"""),

    # Cell 7: Text-to-Pandas với schema context
    cell("""\
def answer(question, df, schema, chat_cfg, sheet_name):
    rev_col  = schema["key_columns"]["main_revenue"]
    desc = schema["sheet_description"]
    struct = schema["row_structure"]

    prompt = f\"\"\"Bạn là data analyst. DataFrame `df` chứa dữ liệu từ sheet "{sheet_name}".

Mô tả: {desc}
Cấu trúc: {struct}

Cột quan trọng:
- __store__: tên cửa hàng (propagated từ bold parent rows)
- __category__: nhóm sản phẩm (Bánh, Bánh Noel, Quà Tết, Kem Givral, Bánh trung thu) — None nếu là parent row
- Doanh thu kỳ này: "{rev_col}"
- Các cột khác: {list(df.columns[2:10])}

Columns: {list(df.columns[:15])}

Sample (3 rows):
{df[["__store__","__category__",rev_col]].dropna(subset=["__category__"]).head(3).to_string(index=False)}

Câu hỏi: {question}

Quy tắc:
- Chỉ trả về Python code thuần, KHÔNG markdown, KHÔNG giải thích
- Gán kết quả vào biến `result`
- Filter tên cửa hàng: dùng df['__store__'].str.contains('...', case=False, na=False)
- Chỉ lấy sub-rows (category != None) khi cần breakdown, parent rows khi cần tổng
\"\"\"
    import re, traceback
    kw   = {{k: v for k, v in chat_cfg.items() if k != "model" and v is not None}}
    resp = litellm.completion(
        model=chat_cfg["model"],
        messages=[{{"role":"user","content":prompt}}],
        temperature=0, **kw)
    code = resp.choices[0].message.content.strip()
    code = re.sub(r"^```python\\n?","",code); code = re.sub(r"^```\\n?","",code); code = re.sub(r"\\n?```$","",code)
    print(f"Code:\\n{{code}}\\n")
    local = {{"df": df, "pd": pd, "np": np}}
    try:
        exec(code, {{}}, local)
        return local.get("result")
    except Exception:
        return traceback.format_exc()

q = "Doanh thu thực hiện tháng 1/2026 của cửa hàng 98 Nguyễn Văn Trỗi là bao nhiêu?"
result = answer(q, df_full, schema, chat_cfg, "DT 01.2026")
print("=== RESULT ===")
print(result)
"""),

    # Cell 8: More test questions
    cell("""\
q2 = "Top 3 cửa hàng có doanh thu Bánh cao nhất tháng 1?"
result2 = answer(q2, df_full, schema, chat_cfg, "DT 01.2026")
print("=== RESULT ===")
print(result2)
"""),

    cell("""\
q3 = "Tổng doanh thu toàn công ty (chỉ tính nhóm Bánh) tháng 1/2026?"
result3 = answer(q3, df_full, schema, chat_cfg, "DT 01.2026")
print("=== RESULT ===")
print(result3)
"""),
]

nb = {
    "cells": cells,
    "metadata": {
        "kernelspec": {"display_name":"Python 3","language":"python","name":"python3"},
        "language_info": {"name":"python","version":"3.12.0"},
    },
    "nbformat": 4,
    "nbformat_minor": 5,
}

with open("test_llm_schema.ipynb", "w", encoding="utf-8") as f:
    json.dump(nb, f, ensure_ascii=False, indent=1)
print("Written: test_llm_schema.ipynb")
